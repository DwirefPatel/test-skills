"""
Folder intake: point it at a client's documents and it prepares the Drake data entry.

  python run.py intake "Clients/Rivera"            (folder holds client.json + documents)
  python run.py intake "Clients"                    (one sub-folder per client, each with client.json)
  python run.py intake <folder> --ssn 123-45-6789 --name "Sam Rivera"   (no client.json)

For every file it:
  1. reads the text (PDF via pdftotext/pypdf, .txt, or a keyed brokerage .json package);
  2. classifies it: brokerage consolidated 1099, 1099-DIV/INT/OID/MISC, W-2, or unrecognized;
  3. checks it belongs to this client (recipient / employee SSN) and this tax year;
  4. 1099s: builds the tied-out brokerage workpaper and stages Schedule D (8949) and
     Schedule B (DIV/INT/OID/99M) rows - ONLY if every tie-out passes;
     W-2s: parses the boxes and stages a W2 row (flagged for review against the PDF);
  5. writes the rows to the input workbook (backed up first) and logs everything.

Nothing is typed into Drake here. Then: `validate`, `enter` (or run intake with --enter).

A file can carry preparer-keyed corrections in a sidecar '<file name>.keyed.json'
(brokerage: box totals / payer / header fields; W-2: W2 tab field names).
"""
import json
import re
import shutil
import sys
import time
from collections import OrderedDict
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .brokerage_rows import build_rows
from .inputs import COMMON, EXAMPLE_ROW, NAME_ROW, write_form_sheet
from .util import digits, mask, sha256_file
from .w2_parser import is_w2, parse_w2

PREPARED = "Prepared"
REFUSED = "Refused - did not tie"
NEEDS = "Needs keyed info"
NOT_AUTOMATED = "Recognized - not automated"
UNRECOGNIZED = "Unrecognized"
DUPLICATE = "Duplicate file - skipped"
SUPERSEDED = "Superseded by CORRECTED"
ALREADY = "Already prepared - skipped"
WRONG_CLIENT = "Refused - another person's document"
WRONG_YEAR = "Refused - wrong tax year"
DONE_STATUSES = {PREPARED, DUPLICATE, SUPERSEDED, ALREADY}

TEXT_TYPES = {".pdf", ".txt"}
IMAGE_TYPES = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".heic", ".bmp", ".gif"}
FORM_PATTERNS = OrderedDict([("B", r"1099-B\b"), ("DA", r"1099-DA\b"), ("DIV", r"1099-DIV"),
                             ("INT", r"1099-INT"), ("OID", r"1099-OID"), ("MISC", r"1099-MISC")])
PAYER_WORDS = re.compile(r"\b(LLC|L\.L\.C\.|INC|INCORPORATED|CORP|CORPORATION|BANK|N\.A\.|SECURITIES|"
                         r"BROKERAGE|FINANCIAL|INVESTMENTS?|TRUST|COMPANY|& CO\.?|CLEARING|SERVICES)\b",
                         re.I)


@dataclass
class Doc:
    path: Path
    sha: str
    category: str = ""
    status: str = ""
    detail: str = ""
    kind: str = ""                     # "1099" | "W2" | ""
    pages: list = field(default_factory=list)
    package: object = None
    w2: list = field(default_factory=list)
    sidecar: object = None
    ts: str = "T"
    rows: dict = field(default_factory=dict)
    manual: list = field(default_factory=list)
    review: list = field(default_factory=list)
    fix_first: list = field(default_factory=list)
    workpaper: str = ""
    workpaper_sha: str = ""
    tie: str = ""
    source_tag: str = ""

    @property
    def name(self):
        return self.path.name

    def set(self, status, detail=""):
        self.status, self.detail = status, detail
        return self


# ----------------------------------------------------------------------------- helpers
def load_brokerage(base):
    try:
        import brokerage_1099
    except ImportError:
        src = Path(base).resolve().parent / "src"
        if not (src / "brokerage_1099").exists():
            raise SystemExit("The brokerage_1099 package is not installed. From the repository root run:\n"
                             "  pip install -e .")
        sys.path.insert(0, str(src))
        import brokerage_1099
    import brokerage_1099.parsers  # noqa: F401  (registers parsers)
    return brokerage_1099


def read_pages(path, b1099):
    if path.suffix.lower() == ".txt":
        return path.read_text(encoding="utf-8", errors="replace").split("\f")
    try:
        return b1099.parsers.extract_pages(path)
    except b1099.parsers.ParseError as exc:
        try:
            from pypdf import PdfReader
        except ImportError:
            raise b1099.parsers.ParseError(f"{exc}; or install pypdf (pip install pypdf)")
        return [p.extract_text(extraction_mode="layout") or "" for p in PdfReader(str(path)).pages]


def classify(pages):
    text = "\n".join(pages)
    up = text.upper()
    if is_w2(text) and not re.search(r"1099-[A-Z]", up):
        return "W-2", "W2"
    has = {k: bool(re.search(p, up)) for k, p in FORM_PATTERNS.items()}
    income = [k for k in ("DIV", "INT", "OID", "MISC") if has[k]]
    if has["B"] or has["DA"] or len(income) > 1:
        return "Brokerage consolidated 1099", "1099"
    if len(income) == 1:
        return f"1099-{income[0]}", "1099"
    if re.search(r"\b1098\b|1095-|SCHEDULE K-1|1099-R\b|1099-NEC|SSA-1099|1099-G\b", up):
        return "Other tax document (not automated)", ""
    return "", ""


def guess_payer_name(pages):
    for line in (pages[0].splitlines() if pages else [])[:20]:
        for seg in re.split(r"\s{3,}", line.strip()):
            if PAYER_WORDS.search(seg) and not re.search(r"\d{2}-\d{7}|1099|FORM\b", seg, re.I) \
                    and 3 < len(seg) <= 60:
                return seg.strip(" ,")
    return ""


def load_clients(folder, args):
    folder = Path(folder)
    if not folder.is_dir():
        raise SystemExit(f"Folder not found: {folder}")

    def spec(path, cj):
        c = json.loads(cj.read_text(encoding="utf-8")) if cj and cj.exists() else {}
        if getattr(args, "ssn", None):
            c["ssn"] = args.ssn
        if getattr(args, "name", None):
            c["name"] = args.name
        if getattr(args, "client_id", None):
            c["client_id"] = args.client_id
        if getattr(args, "ts", None):
            c["ts"] = args.ts
        if len(digits(c.get("ssn"))) != 9:
            raise SystemExit(f"{path}: client SSN missing or not 9 digits. Put it in {path / 'client.json'} "
                             f'as {{"ssn": "123-45-6789", "name": "Client Name"}} or pass --ssn.')
        c.setdefault("name", "")
        c.setdefault("client_id", "")
        c["ts"] = (c.get("ts") or "T").upper()
        c["key"] = c["client_id"] or mask(c["ssn"])
        return c

    if getattr(args, "ssn", None) or (folder / "client.json").exists():
        return [(folder, spec(folder, folder / "client.json"))]
    subs = [d for d in sorted(folder.iterdir()) if d.is_dir() and (d / "client.json").exists()]
    if not subs:
        raise SystemExit(f"No client.json in {folder} or its sub-folders, and no --ssn given.")
    return [(d, spec(d, d / "client.json")) for d in subs]


def skip_file(p):
    n = p.name.lower()
    return (n == "client.json" or n.endswith(".keyed.json") or n.startswith((".", "~$"))
            or n in ("thumbs.db", "desktop.ini"))


def which_person(found, client):
    """'T', 'S', or None (someone else) for an SSN / masked SSN found on a document; '' if none."""
    d = digits(found)
    if not d:
        return ""
    for who, ssn in (("T", client["ssn"]), ("S", client.get("spouse_ssn", ""))):
        s = digits(ssn)
        if s and (d == s or (len(d) == 4 and s.endswith(d))):
            return who
    return None


# ----------------------------------------------------------------------------- per document
def read_doc(doc, b1099, client, settings):
    """Classify and parse. Leaves doc.status empty when ready for the brokerage/W-2 step."""
    p, suffix = doc.path, doc.path.suffix.lower()
    side = p.with_name(p.name + ".keyed.json")
    if side.exists():
        doc.sidecar = json.loads(side.read_text(encoding="utf-8"))
    if suffix in IMAGE_TYPES:
        doc.category = "Image"
        return doc.set(NEEDS, "Image file: scan it to a searchable PDF (OCR) and drop that in instead")
    if suffix == ".json":
        data = json.loads(p.read_text(encoding="utf-8"))
        if not any(k in data for k in ("lots", "form_totals", "income_lines")):
            doc.category = "JSON (unknown)"
            return doc.set(UNRECOGNIZED, "JSON file is not a brokerage 1099 package")
        doc.category, doc.kind = "Brokerage 1099 (keyed JSON)", "1099"
        data.setdefault("client", client["name"] or client["key"])
        doc.package = b1099.package_from_dict(data)
    elif suffix in TEXT_TYPES:
        try:
            doc.pages = read_pages(p, b1099)
        except b1099.parsers.ParseError as exc:
            doc.category = "PDF"
            return doc.set(NEEDS, f"Could not read the text: {exc}")
        if not "".join(doc.pages).strip():
            doc.category = "PDF (no text)"
            return doc.set(NEEDS, "No text layer (scanned image): OCR it, or key it in a .keyed.json")
        doc.category, doc.kind = classify(doc.pages)
        if not doc.category:
            doc.category = "Unknown"
            return doc.set(UNRECOGNIZED, "Not a W-2 or brokerage/1099 statement")
        if not doc.kind:
            return doc.set(NOT_AUTOMATED, "Recognized, but this document type isn't automated yet")
    else:
        doc.category = f"{suffix or 'no extension'} file"
        return doc.set(UNRECOGNIZED, "File type not supported (use PDF, TXT or a brokerage JSON)")

    if doc.kind == "W2":
        return read_w2(doc, client)
    return read_1099(doc, b1099, client, settings)


def read_w2(doc, client):
    parsed = parse_w2(doc.pages)
    keyed = doc.sidecar or []
    keyed = keyed if isinstance(keyed, list) else [keyed]
    if not parsed and keyed:
        parsed = [({}, []) for _ in keyed]
    if not parsed:
        return doc.set(NEEDS, "Looks like a W-2 but no boxes could be read; key it in a .keyed.json")
    for i, (row, missing) in enumerate(parsed):
        if i < len(keyed):
            row.update({k: str(v) for k, v in keyed[i].items()})
            missing = [m for m in missing if not row.get(m)]
        who = which_person(row.pop("_employee_ssn", ""), client)
        if who is None:
            return doc.set(WRONG_CLIENT, f"Employee SSN on W-2 #{i + 1} is not the client's or spouse's")
        row["ts"] = who or row.get("ts") or client["ts"]
        if not who:
            doc.review.append("Employee SSN not found on the W-2: confirm it belongs to this client")
        if missing:
            doc.review.append(f"W-2 #{i + 1}: could not read {', '.join(missing)} - key them from the PDF")
        doc.w2.append(row)
    doc.review.append("Auto-parsed W-2: compare every box with the PDF before entry (box 12, 13 and 14 "
                      "especially)")
    return doc


def read_1099(doc, b1099, client, settings):
    if doc.package is None:
        broker = (doc.sidecar or {}).get("parser")
        try:
            doc.package = b1099.parsers.parse_text(doc.pages, broker=broker, source_file=doc.name)
        except b1099.parsers.ParseError as exc:
            return doc.set(NEEDS, f"Parser could not map this layout: {exc}")
    pkg = doc.package
    if doc.sidecar:
        pkg = doc.package = b1099.merge_overrides(pkg, doc.sidecar)
    pkg.client = pkg.client or client["name"] or client["key"]
    payer_guess = guess_payer_name(doc.pages) if doc.pages else ""
    if pkg.broker in ("", "Unknown broker"):
        pkg.broker = payer_guess or pkg.broker
    for payer in pkg.payers:
        if payer.name in ("", "Unknown broker"):
            payer.name = payer_guess or pkg.broker
    year = settings.get("tax_year")
    if pkg.tax_year and year and pkg.tax_year != int(year):
        return doc.set(WRONG_YEAR, f"Statement is for {pkg.tax_year}; settings tax_year is {year}")
    if not pkg.tax_year:
        pkg.tax_year = int(year)
        doc.review.append(f"Tax year not found on the statement; assumed {year}")
    who = which_person(pkg.recipient_tin_last4, client)
    if who is None:
        return doc.set(WRONG_CLIENT, f"Recipient TIN {pkg.recipient_tin_last4} is not the client's or spouse's")
    doc.ts = who or client["ts"]
    if not pkg.payers or not pkg.payers[0].ein:
        return doc.set(NEEDS, "Payer EIN not found; add it in a .keyed.json (payers: name, ein)")
    if pkg.payers[0].name in ("", "Unknown broker"):
        return doc.set(NEEDS, "Payer name not found; add it in a .keyed.json (payers: name, ein)")
    if not (pkg.lots or pkg.income_lines or pkg.form_totals or pkg.other_items):
        return doc.set(NEEDS, "Nothing could be read from this layout: add a broker parser, or key the "
                              "statement as JSON")
    doc.source_tag = (f"1099 {pkg.payers[0].name[:30]} EIN {pkg.payers[0].ein} acct "
                      f"{pkg.payers[0].account_last4 or '-'} TY{pkg.tax_year}")
    return doc


def resolve_duplicates(docs):
    by_sha = {}
    for d in docs:
        if d.sha in by_sha and not d.status:
            d.set(DUPLICATE, f"Same file as {by_sha[d.sha].name}")
        by_sha.setdefault(d.sha, d)
    groups = OrderedDict()
    for d in docs:
        if d.kind == "1099" and not d.status:
            groups.setdefault(d.source_tag, []).append(d)
    for tag, ds in groups.items():
        if len(ds) < 2:
            continue
        corrected = [d for d in ds if d.package.corrected]
        if len(corrected) == 1:
            for d in ds:
                if d is not corrected[0]:
                    d.set(SUPERSEDED, f"{corrected[0].name} is the CORRECTED statement for this account")
        else:
            for d in ds:
                d.set(NEEDS, "More than one statement for the same payer and account: remove the extra "
                             "one (or mark which is CORRECTED)")


def prepare_1099(doc, b1099, client, forms, out_dir, recalc, settings):
    pkg = doc.package
    result = b1099.process_package(pkg, out_dir, filename=f"{doc.path.stem}_1099_Summary.xlsx",
                                   recalc=recalc)
    doc.workpaper = str(result.workbook_path)
    doc.workpaper_sha = sha256_file(result.workbook_path)
    missing_totals = [w for w in result.validation.warnings if "no broker summary total" in w]
    problems = [f"{c.name}: diff {c.diff}" for c in result.validation.failures]
    if result.verify is not None:
        if result.verify.recalc.error:
            problems.append(result.verify.recalc.error)
        problems += [f"{c.location} = {c.value}" for c in result.verify.failures]
        problems += result.verify.recalc.formula_errors[:5]
        doc.tie = "ALL TIED" if not problems else "CHECK"
    else:
        doc.tie = "Python tie-outs only (workpaper not recalculated)" if not problems else "CHECK"
        if settings.get("import_1099", {}).get("require_recalc"):
            problems.append("LibreOffice is required (settings import_1099.require_recalc) but not available")
    if problems:
        return doc.set(REFUSED, "; ".join(problems[:8]))
    if missing_totals:
        return doc.set(NEEDS, "Broker summary totals missing (cannot tie the lots): "
                              + "; ".join(missing_totals) + " - key box_totals in a .keyed.json")
    doc.rows, doc.manual = build_rows(result, {**client, "ts": doc.ts}, doc.source_tag, forms)
    return doc.set(PREPARED, result.headline)


def prepare_w2(doc, client):
    rows = []
    for w in doc.w2:
        rows.append({"client_ssn": client["ssn"], "client_id": client["client_id"],
                     "client_name": client["name"], **w, "source_doc": f"W-2 {doc.name}"})
    doc.rows = {"W2": rows}
    return doc.set(PREPARED, f"{len(rows)} W-2(s) parsed")


# ----------------------------------------------------------------------------- workbook
def _row_key(code, row):
    if code == "8949":
        return row.get("lot_ref", "")
    if code == "W2":
        return digits(row.get("employer_ein", ""))
    return digits(row.get("payer_ein", ""))


def _belongs(code, row, client_ssn, doc_rows, tag):
    if digits(row.get("client_ssn")) != client_ssn:
        return False
    if code == "W2":
        return _row_key(code, row) in {_row_key(code, r) for r in doc_rows}
    return row.get("source_doc", "") == tag


def _check_rows(doc, code, rows, forms, settings):
    for r in rows:
        errors, _ = forms[code].validate(r, settings)
        ref = r.get("lot_ref") or r.get("payer_name") or r.get("employer_name") or ""
        doc.fix_first += [f"{code} {ref}: {e}" for e in errors]


def write_rows(input_path, forms, client, docs, replace, settings):
    """Stage each prepared doc's rows, then validate what ended up in the workbook.
    Returns {code: rows_written}."""
    staged = [d for d in docs if d.status == PREPARED and d.rows]
    if not staged:
        return {}
    wb = load_workbook(input_path)
    written = {}
    ssn = digits(client["ssn"])
    for doc in staged:
        plan = []
        for code, new_rows in doc.rows.items():
            form = forms[code]
            if form.sheet not in wb.sheetnames:
                write_form_sheet(wb, form)
            ws = wb[form.sheet]
            names = [c.value for c in ws[NAME_ROW]]
            for f in COMMON + list(form.fields):          # columns added by newer versions
                if f.name not in names:
                    names.append(f.name)
                    ws.cell(NAME_ROW, len(names), f.name).font = Font(name="Arial", size=8, color="808080")
                    ws.cell(1, len(names), f.label).font = Font(name="Arial", bold=True, color="FFFFFF", size=9)
                    ws.cell(1, len(names)).fill = PatternFill("solid", fgColor="1F4E78")
            col = {n: j for j, n in enumerate(names, 1) if n}
            existing = []
            for i in range(EXAMPLE_ROW + 1, ws.max_row + 1):
                vals = {n: ws.cell(i, j).value for n, j in col.items()}
                if any(v not in (None, "") for v in vals.values()):
                    existing.append((i, {k: "" if v is None else str(v) for k, v in vals.items()}))
            mine = [(i, r) for i, r in existing if _belongs(code, r, ssn, new_rows, doc.source_tag)]
            plan.append((code, ws, col, new_rows, existing, mine))
        if any(mine for *_, mine in plan) and not replace:
            doc.set(ALREADY, "Rows for this document are already in the input workbook (use --replace to "
                             "refresh them; values you typed into blank cells are kept)")
            for code, *_, mine in plan:
                _check_rows(doc, code, [r for _, r in mine], forms, settings)
            continue
        for code, ws, col, new_rows, existing, mine in plan:
            old = {_row_key(code, r): r for _, r in mine}
            for i, _ in sorted(mine, reverse=True):
                ws.delete_rows(i)
            kept = [i for i, _ in existing if i not in {m for m, _ in mine}]
            nxt = max([EXAMPLE_ROW] + [i - sum(1 for m, _ in mine if m < i) for i in kept]) + 1
            for row in new_rows:
                prev = old.get(_row_key(code, row), {})
                for k, v in prev.items():               # keep what the preparer filled in
                    if v and not str(row.get(k, "")).strip():
                        row[k] = v
                for name, j in col.items():
                    c = ws.cell(nxt, j, str(row.get(name, "") or "") or None)
                    c.number_format = "@"
                    c.font = Font(name="Arial", size=10)
                nxt += 1
            written[code] = written.get(code, 0) + len(new_rows)
            _check_rows(doc, code, new_rows, forms, settings)
    try:
        wb.save(input_path)
    except PermissionError:
        raise SystemExit(f"{input_path.name} is open in Excel. Close it and run intake again.")
    return written


# ----------------------------------------------------------------------------- report
def write_report(path, all_docs):
    wb = Workbook()
    ws = wb.active
    ws.title = "Documents"
    head_fill = PatternFill("solid", fgColor="1F4E78")
    fills = {"ok": PatternFill("solid", fgColor="E2EFDA"), "bad": PatternFill("solid", fgColor="F8CBAD"),
             "mid": PatternFill("solid", fgColor="FFE699")}

    def table(ws, heads, widths, rows, status_col=None):
        for j, (h, w) in enumerate(zip(heads, widths), 1):
            c = ws.cell(1, j, h)
            c.font = Font(name="Arial", bold=True, color="FFFFFF")
            c.fill = head_fill
            ws.column_dimensions[c.column_letter].width = w
        for i, r in enumerate(rows, 2):
            for j, v in enumerate(r, 1):
                c = ws.cell(i, j, v)
                c.font = Font(name="Arial", size=10)
                c.alignment = Alignment(wrap_text=True, vertical="top")
            if status_col:
                s = r[status_col - 1]
                ws.cell(i, status_col).fill = fills["ok" if s in DONE_STATUSES else
                                                      "mid" if s in (NOT_AUTOMATED, UNRECOGNIZED) else "bad"]
        ws.freeze_panes = "A2"

    docs_rows, manual_rows, fix_rows = [], [], []
    for client, docs in all_docs:
        for d in docs:
            staged = ", ".join(f"{k}: {len(v)}" for k, v in d.rows.items()) if d.status == PREPARED else ""
            docs_rows.append([client["key"], client["name"], d.name, d.category, d.status, d.detail,
                              staged, d.tie, Path(d.workpaper).name if d.workpaper else "",
                              "\n".join(d.review), d.sha[:16]])
            for item, amt, dest in d.manual:
                manual_rows.append([client["key"], d.name, item, None if amt is None else float(amt), dest])
            for e in d.fix_first:
                fix_rows.append([client["key"], d.name, e])
    table(ws, ["Client", "Name", "File", "Category", "Status", "Detail", "Rows staged", "Tie-out",
               "Workpaper", "Review before entry", "SHA-256"],
          [14, 18, 32, 26, 26, 60, 22, 16, 34, 50, 18], docs_rows, status_col=5)
    table(wb.create_sheet("Fix before entry"), ["Client", "File", "Validation error (validate will stop "
                                                "until fixed in the input workbook)"],
          [14, 32, 110], fix_rows)
    m = wb.create_sheet("Enter manually")
    table(m, ["Client", "File", "Item", "Amount", "Where it goes"], [14, 32, 46, 14, 60], manual_rows)
    for i in range(2, len(manual_rows) + 2):
        m.cell(i, 4).number_format = '#,##0.00;(#,##0.00);"-"'
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ----------------------------------------------------------------------------- driver
def run_intake(session, folder):
    """Returns [(client, docs)]."""
    args = session.args
    b1099 = load_brokerage(session.base)
    from brokerage_1099.recalc import soffice_binary

    recalc = soffice_binary() is not None and not getattr(args, "no_recalc", False)
    need = ["W2", "8949", "DIV", "INT", "OID", "99M"]
    missing = [c for c in need if c not in session.forms]
    if missing:
        from forms import load_forms
        session.forms.update(load_forms(missing, session.base))
    forms = session.forms
    input_path = session.paths["input"]
    backup = input_path.with_name(f"{input_path.stem}_backup_{time.strftime('%Y%m%d_%H%M%S')}{input_path.suffix}")
    shutil.copy2(input_path, backup)
    print(f"Backed up the input workbook to {backup.name}")
    if not recalc:
        print("Note: LibreOffice not available (or --no-recalc): workpapers are gated on the Python "
              "tie-outs only.")

    results = []
    for cfolder, client in load_clients(folder, args):
        print(f"\n=== {client['key']}  {client['name']}  ({cfolder}) ===")
        docs = []
        for p in sorted(x for x in cfolder.iterdir() if x.is_file() and not skip_file(x)):
            docs.append(read_doc(Doc(path=p, sha=sha256_file(p)), b1099, client, session.settings))
        resolve_duplicates(docs)
        out_dir = session.paths.get("workpapers", session.base / "workpapers") / \
            re.sub(r"[^A-Za-z0-9_-]+", "_", f"{client['key'][-4:]}_{client['name']}").strip("_")
        for d in docs:
            if d.status:
                continue
            if d.kind == "W2":
                prepare_w2(d, client)
            else:
                prepare_1099(d, b1099, client, forms, out_dir, recalc, session.settings)
        written = write_rows(input_path, forms, client, docs, getattr(args, "replace", False),
                             session.settings)
        for d in docs:
            rel = d.workpaper
            try:
                rel = str(Path(d.workpaper).resolve().relative_to(session.base)) if d.workpaper else ""
            except ValueError:
                pass
            session.log("intake_document", client=client["key"], client_name=client["name"], file=d.name,
                        sha256=d.sha, category=d.category, status=d.status, detail=d.detail[:500],
                        rows={k: len(v) for k, v in d.rows.items()} if d.status == PREPARED else {},
                        tie=d.tie, workpaper=rel, workpaper_sha256=d.workpaper_sha, review=d.review,
                        fix_first=d.fix_first[:50],
                        manual=[[i, None if a is None else float(a), w] for i, a, w in d.manual])
            mark = "OK " if d.status in DONE_STATUSES else "-- "
            print(f"  {mark}{d.name[:40]:40s} {d.category[:28]:28s} {d.status}")
            if d.status not in DONE_STATUSES and d.detail:
                print(f"       {d.detail[:200]}")
        if written:
            print("  Staged in the input workbook: " + ", ".join(f"{k} {v} row(s)" for k, v in written.items()))
        fixes = sum(len(d.fix_first) for d in docs)
        if fixes:
            print(f"  {fixes} item(s) to fix in the input workbook before 'enter' (see the intake report).")
        results.append((client, docs))
    return results
