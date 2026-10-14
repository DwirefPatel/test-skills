"""Input workbook: one tab per form, one row per source document."""
from collections import OrderedDict

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

from .fields import Field
from .util import digits, mask

COMMON = [
    Field("client_ssn", "Client SSN (primary taxpayer on the return)", "text", required=True),
    Field("client_id", "Client ID (optional; needed if two clients share SSN last 4)", "text"),
    Field("client_name", "Client name (optional; names the tracker tab)", "text"),
]
HEADER_ROW, NAME_ROW, EXAMPLE_ROW = 1, 2, 3


def load_rows(path, form):
    try:
        df = pd.read_excel(path, sheet_name=form.sheet, header=NAME_ROW - 1, dtype=str).fillna("")
    except ValueError:
        return []   # tab not in workbook
    rows = []
    for i, rec in enumerate(df.to_dict("records")):
        r = {k: str(v).strip() for k, v in rec.items()}
        if not r.get("client_ssn") or r["client_ssn"].upper().startswith("EXAMPLE"):
            continue
        r["_row"] = i + NAME_ROW + 1
        r["_sheet"] = form.sheet
        rows.append(r)
    return rows


def group_clients(rows_by_form):
    """-> (clients, errors). clients: OrderedDict ssn_digits -> {key, name, forms: {code: [rows]}}"""
    clients, errors, key_owner = OrderedDict(), [], {}
    for code, rows in rows_by_form.items():
        for r in rows:
            ssn = digits(r["client_ssn"])
            if len(ssn) != 9:
                errors.append(f"{r['_sheet']} row {r['_row']}: client SSN must be 9 digits")
                continue
            key = r.get("client_id") or mask(ssn)
            if key_owner.setdefault(key, ssn) != ssn:
                errors.append(f"{r['_sheet']} row {r['_row']}: two different clients both map to '{key}'. "
                              f"Fill in client_id for them.")
                continue
            c = clients.setdefault(ssn, {"key": key, "name": "", "forms": OrderedDict()})
            if r.get("client_name") and not c["name"]:
                c["name"] = r["client_name"]
            c["forms"].setdefault(code, []).append(r)
    return clients, errors


def make_template(path, forms, examples=None):
    examples = examples or {}
    wb = Workbook()
    info = wb.active
    info.title = "Instructions"
    lines = [
        ("Client input workbook", True),
        ("One tab per form. One row per source document (e.g. one row per W-2).", False),
        ("Row 1 = field description, Row 2 = field name used by the script (don't edit rows 1-2).", False),
        ("Row 3 = EXAMPLE row showing the format; the script ignores it. Enter data from row 4 down.", False),
        ("All cells are Text format so leading zeros in EINs, ZIPs and IDs are kept.", False),
        ("Money: plain numbers; $ and commas are fine. Checkboxes: X or blank.", False),
        ("Rows with the same client SSN go into the same Drake return.", False),
        ("Run 'python run.py validate' after filling in; fix every ERROR before entering.", False),
        ("SECURITY: this file contains taxpayer data. Keep it on the secured PC only.", True),
    ]
    for i, (t, bold) in enumerate(lines, 1):
        c = info.cell(i, 1, t)
        c.font = Font(name="Arial", bold=bold, size=12 if i == 1 else 10)
    info.column_dimensions["A"].width = 110

    for form in forms.values():
        write_form_sheet(wb, form, examples.get(form.code, {}))
    wb.save(path)


def write_form_sheet(wb, form, example=None):
    """Add one form's tab (header, field names, example row, text-formatted input area)."""
    example = example or {}
    head_fill = PatternFill("solid", fgColor="1F4E78")
    yellow = PatternFill("solid", fgColor="FFF2CC")
    ws = wb.create_sheet(form.sheet)
    cols = COMMON + list(form.fields)
    for j, f in enumerate(cols, 1):
        h = ws.cell(HEADER_ROW, j, f.label)
        h.font = Font(name="Arial", bold=True, color="FFFFFF", size=9)
        h.fill = head_fill
        h.alignment = Alignment(wrap_text=True, vertical="top")
        n = ws.cell(NAME_ROW, j, f.name)
        n.font = Font(name="Arial", size=8, color="808080")
        if f.help:
            h.comment = Comment(f.help, "Template")
        ex = ws.cell(EXAMPLE_ROW, j, example.get(f.name, ""))
        ex.font = Font(name="Arial", italic=True, color="808080")
        ex.number_format = "@"
        for i in range(EXAMPLE_ROW + 1, EXAMPLE_ROW + 301):
            c = ws.cell(i, j)
            c.number_format = "@"
            c.font = Font(name="Arial", size=10)
            if i <= EXAMPLE_ROW + 25:
                c.fill = yellow
        ws.column_dimensions[h.column_letter].width = max(13, min(30, len(f.name) + 3))
    ws.cell(EXAMPLE_ROW, 1).value = "EXAMPLE-" + str(ws.cell(EXAMPLE_ROW, 1).value or "")
    ws.row_dimensions[HEADER_ROW].height = 48
    ws.freeze_panes = ws.cell(EXAMPLE_ROW, 4)
    return ws
