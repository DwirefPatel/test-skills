"""
Builds tracking/Entry_Tracker.xlsx from the audit log.

Sheets
  Summary          one line per client per form, status formulas, preparer/reviewer sign-off
  Exceptions       everything that needs a human: open mismatches, incomplete entries,
                   failed/unconfirmed tie-outs, accepted validation warnings, fields not read back
  Corrections      every correction: before / after / source / reason / who / when
  Tie-Out          script totals vs the return, and whether the preparer confirmed they agree
  Run Log          every run: who, when, mode, input file + SHA-256, outcome
  Audit Integrity  result of checking the hash chain of the audit log
  <client tabs>    one section per form (page break before each), one block per document:
                   Drake Screen Field | Source Document | Entered in Drake | Verified |
                   Last Action | Reviewer | Notes
"""
import os
import re
import time
from collections import OrderedDict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.pagebreak import Break

from .state import State

FONT = "Arial"
NAVY = PatternFill("solid", fgColor="1F4E78")
LIGHT = PatternFill("solid", fgColor="DDEBF7")
GREY = PatternFill("solid", fgColor="F2F2F2")
RED = PatternFill("solid", fgColor="F8CBAD")
AMBER = PatternFill("solid", fgColor="FFE699")
GREEN = PatternFill("solid", fgColor="E2EFDA")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
ACTION = {"enter": "Entered", "verify": "Verified", "correct": "Corrected"}


def _f(bold=False, size=10, color="000000", italic=False):
    return Font(name=FONT, bold=bold, size=size, color=color, italic=italic)


def _q(sheet):
    return "'" + sheet.replace("'", "''") + "'"


def _blank(v):
    s = re.sub(r"[\s$,]", "", str(v or ""))
    if s == "":
        return True
    try:
        return float(s) == 0
    except ValueError:
        return False


def _header(ws, row, heads, widths=None):
    for j, t in enumerate(heads, 1):
        c = ws.cell(row, j, t)
        c.font = _f(True, 10, "FFFFFF")
        c.fill = NAVY
        c.alignment = Alignment(wrap_text=True, vertical="center")
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row, j).column_letter].width = w


def _title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = _f(True, 14)
    ws["A2"] = sub
    ws["A2"].font = _f(size=9, color="595959", italic=True)


def _print_setup(ws, last_col, last_row, title_rows="1:2"):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = title_rows
    ws.print_area = f"A1:{last_col}{max(last_row, 3)}"
    ws.oddFooter.left.text = "Confidential - contains taxpayer information"
    ws.oddFooter.right.text = "Page &P of &N"


def _sheet_name(key, name, used):
    base = re.sub(r"[\[\]:*?/\\']", "", f"{key[-4:]} {name}".strip())[:31] or key[-4:]
    out, n = base, 2
    while out.lower() in used:
        out = f"{base[:28]}~{n}"
        n += 1
    used.add(out.lower())
    return out


def build_tracker(session):
    events = session.audit.events()
    if not events:
        return None
    state = State(events)
    ok, n_events, problem = session.audit.verify()
    forms = session.forms
    show_blank = session.settings.get("tracker", {}).get("show_blank_fields", False)
    out_path = Path(session.paths["tracker"])
    out_path.parent.mkdir(parents=True, exist_ok=True)

    clients = OrderedDict()
    for (ck, fcode, pos), inst in state.instances.items():
        c = clients.setdefault(ck, {"name": "", "forms": OrderedDict()})
        c["name"] = inst["client_name"] or c["name"]
        c["forms"].setdefault(fcode, []).append(inst)

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    exc_ws = wb.create_sheet("Exceptions")
    cor_ws = wb.create_sheet("Corrections")
    tie_ws = wb.create_sheet("Tie-Out")
    run_ws = wb.create_sheet("Run Log")
    aud_ws = wb.create_sheet("Audit Integrity")
    used = {s.lower() for s in wb.sheetnames}

    exceptions = []
    summary_rows = []
    for ck, c in clients.items():
        ws = wb.create_sheet(_sheet_name(ck, c["name"], used))
        _client_sheet(ws, ck, c, forms, show_blank, out_path, session.base, exceptions, summary_rows)

    latest_ties = _tie_out_sheet(tie_ws, state, exceptions)
    for ck, c in clients.items():
        for fcode in c["forms"]:
            if (ck, fcode) not in latest_ties:
                exceptions.append(("Medium", ck, fcode, "", "Tie-out", "No tie-out recorded yet "
                                   "(run verify for this client to do it)", "", "", "", ""))
    _corrections_sheet(cor_ws, state)
    _validation_exceptions(state, exceptions)
    _exceptions_sheet(exc_ws, exceptions)
    _run_log_sheet(run_ws, state)
    _integrity_sheet(aud_ws, ok, n_events, problem, events, session)
    _summary_sheet(summary, summary_rows, state, latest_ties)

    try:
        wb.save(out_path)
        return out_path
    except PermissionError:
        alt = out_path.with_name(f"{out_path.stem}_{time.strftime('%Y%m%d_%H%M%S')}.xlsx")
        wb.save(alt)
        print(f"  {out_path.name} is open in Excel; saved {alt.name} instead.")
        return alt


# ---------------------------------------------------------------------- client tabs
def _client_sheet(ws, ck, c, forms, show_blank, out_path, base, exceptions, summary_rows):
    widths = (42, 22, 22, 15, 24, 10, 44, 6, 9, 20)
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, j).column_letter].width = w
    for col in "HIJ":
        ws.column_dimensions[col].hidden = True   # helpers for Summary formulas
    _title(ws, f"{c['name'] or 'Client'} ({ck}) - Drake Entry Tracker",
           "Source Document = input workbook value.  Entered in Drake = value read back from the Drake screen "
           "(typed value where the field can't be read back).  Reviewer: initial each line or each document.")

    order = [f for f in forms if f in c["forms"]] + [f for f in c["forms"] if f not in forms]
    r, first = 4, True
    for fcode in order:
        form = forms.get(fcode)
        title = form.title if form else f"Form {fcode}"
        inst_label = form.instance_label if form else fcode
        src_label = form.source_label if form else fcode
        if not first:
            ws.row_breaks.append(Break(id=r - 1))
        first = False
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        h = ws.cell(r, 1, title.upper())
        h.font = _f(True, 12, "FFFFFF")
        for j in range(1, 8):
            ws.cell(r, j).fill = NAVY
        ws.row_dimensions[r].height = 22
        r += 2

        for inst in c["forms"][fcode]:
            status = State.status(inst)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
            hdr = ws.cell(r, 1, f"{inst_label} #{inst['pos']}  -  {inst['desc']}  -  entered {inst['entered_ts']} "
                                f"by {inst['entered_by']}  -  {status}")
            hdr.font = _f(True, 11)
            fill = RED if status in ("Open mismatch", "Incomplete") else LIGHT
            for j in range(1, 8):
                ws.cell(r, j).fill = fill
            ws.cell(r, 8, fcode)
            ws.cell(r, 9, "INSTANCE")
            ws.cell(r, 10, status)
            r += 1
            if inst["screenshot"]:
                shot = Path(base) / inst["screenshot"]
                rel = os.path.relpath(shot, out_path.parent)
                s = ws.cell(r, 1, f"Evidence: screenshot {Path(inst['screenshot']).name}  "
                                  f"(SHA-256 {inst['screenshot_sha'][:16]}...)")
                s.hyperlink = rel.replace("\\", "/")
                s.font = _f(size=9, color="0563C1", italic=True)
                r += 1
            if inst["last_checked"]:
                ws.cell(r, 1, f"Last verified/corrected: {inst['last_checked']}").font = _f(size=9, italic=True)
                r += 1

            heads = ["Drake Screen Field", f"Source Document ({src_label})", "Entered in Drake", "Verified",
                     "Last Action", "Reviewer", "Notes"]
            for j, t in enumerate(heads, 1):
                x = ws.cell(r, j, t)
                x.font = _f(True)
                x.fill = GREY
            r += 1

            field_order = [n for n in (form.field_order if form else inst["fields"]) if n != "__SKIP__"]
            for name in field_order + [n for n in inst["fields"] if n not in field_order]:
                fr = inst["fields"].get(name)
                if not fr:
                    continue
                if (not show_blank and fr["result"] in ("OK", "Not read back")
                        and _blank(fr["source"]) and _blank(fr["drake"])):
                    continue
                last = f"{ACTION.get(fr.get('action'), '')} {fr.get('ts', '')} {fr.get('by', '')}".strip()
                vals = [fr["label"], fr["source"], fr["drake"], fr["result"], last, "", fr.get("note", "")]
                for j, v in enumerate(vals, 1):
                    x = ws.cell(r, j, "" if v is None else str(v))
                    x.font = _f(size=10)
                    x.border = THIN
                    x.number_format = "@"
                    if j in (2, 3):
                        x.alignment = Alignment(horizontal="right")
                    if j == 7:
                        x.alignment = Alignment(wrap_text=True)
                ws.cell(r, 6).fill = YELLOW
                ws.cell(r, 8, fcode)
                ws.cell(r, 9, "FIELD")
                if fr["result"] == "MISMATCH":
                    exceptions.append(("High", ck, fcode, f"{inst_label} #{inst['pos']}", fr["label"],
                                       "Open mismatch", fr["source"], fr["drake"], fr.get("note", ""), fr.get("ts", "")))
                elif fr["result"] == "Not read back" and not _blank(fr["source"]):
                    exceptions.append(("Info", ck, fcode, f"{inst_label} #{inst['pos']}", fr["label"],
                                       "Not read back - check screenshot", fr["source"], fr["drake"], "",
                                       fr.get("ts", "")))
                r += 1
            if status == "Incomplete":
                exceptions.append(("High", ck, fcode, f"{inst_label} #{inst['pos']}", "(whole document)",
                                   "Entry incomplete - run correct", "", "", "", inst["entered_ts"]))
            r += 1
        summary_rows.append((ws.title, ck, c["name"], fcode, title))
        r += 1

    last = max(r, 5)
    ws.conditional_formatting.add(f"D4:D{last}", CellIsRule(operator="equal", formula=['"MISMATCH"'], fill=RED,
                                                            font=Font(name=FONT, bold=True, color="9C0006")))
    ws.conditional_formatting.add(f"D4:D{last}", CellIsRule(operator="equal", formula=['"OK"'],
                                                            font=Font(name=FONT, color="006100")))
    _print_setup(ws, "G", last)


# ---------------------------------------------------------------------- other sheets
def _tie_out_sheet(ws, state, exceptions):
    _title(ws, "Tie-Out", "Totals from the input workbook to agree to the return in Drake. Latest confirmation "
                          "per client/form/total (full history in the audit log).")
    heads = ["Client", "Name", "Form", "Total", "Amount per Source", "Ties To", "Agrees?", "Confirmed By", "When"]
    _header(ws, 4, heads, (14, 22, 8, 34, 16, 26, 12, 12, 18))
    latest = OrderedDict()
    for t in state.tieouts:
        latest[(t["client"], t["form"], t["label"])] = t
    r = 5
    status = {}
    for (ck, fcode, label), t in latest.items():
        vals = [ck, t.get("client_name", ""), fcode, label, t["amount"], t["tie_to"], t["confirmed"], t["by"], t["ts"]]
        for j, v in enumerate(vals, 1):
            x = ws.cell(r, j, v)
            x.font = _f()
            x.border = THIN
        ws.cell(r, 5).number_format = '#,##0.00;(#,##0.00);"-"'
        status.setdefault((ck, fcode), t["confirmed"])
        if t["confirmed"] == "No":
            status[(ck, fcode)] = "No"
        r += 1
    for (ck, fcode), s in status.items():
        if s == "No":
            exceptions.append(("High", ck, fcode, "", "Tie-out", "Totals did NOT agree to the return", "", "", "",
                               ""))
        elif s != "Yes":
            exceptions.append(("Medium", ck, fcode, "", "Tie-out", "Tie-out not confirmed yet", "", "", "", ""))
    ws.conditional_formatting.add(f"G5:G{max(r, 5)}", CellIsRule(operator="equal", formula=['"No"'], fill=RED))
    ws.conditional_formatting.add(f"G5:G{max(r, 5)}", CellIsRule(operator="equal", formula=['"Yes"'], fill=GREEN))
    ws.freeze_panes = "A5"
    _print_setup(ws, "I", r, "1:4")
    return status


def _corrections_sheet(ws, state):
    _title(ws, "Corrections", "Every change made to Drake after the original entry. 'Before' and 'After' are read "
                              "from the Drake screen.")
    heads = ["When", "Client", "Form", "Document #", "Field", "Before (Drake)", "After (Drake)", "Source Value",
             "Reason", "Result", "By", "Run ID"]
    _header(ws, 4, heads, (18, 14, 7, 11, 36, 18, 18, 18, 26, 16, 7, 16))
    r = 5
    for c in state.corrections:
        vals = [c["ts"], c["client"], c["form"], c["pos"], c["label"], c["before"], c["after"], c["source"],
                c["reason"], c["result"], c["by"], c["run_id"]]
        for j, v in enumerate(vals, 1):
            x = ws.cell(r, j, "" if v is None else str(v))
            x.font = _f()
            x.border = THIN
        r += 1
    ws.conditional_formatting.add(f"J5:J{max(r, 5)}", CellIsRule(operator="equal", formula=['"FAILED"'], fill=RED))
    ws.freeze_panes = "A5"
    _print_setup(ws, "L", r, "1:4")


def _validation_exceptions(state, exceptions):
    enter_runs = [rid for rid, run in state.runs.items() if str(run.get("mode", "")).startswith("enter")
                  and "dry" not in str(run.get("mode", ""))]
    if not enter_runs:
        return
    last = enter_runs[-1]
    for v in state.validations:
        if v["run_id"] == last and v.get("level") == "warning":
            exceptions.append(("Medium", "", "", "", "Input validation", "Warning accepted at entry: " + v["message"],
                               "", "", "", v["ts"]))


def _exceptions_sheet(ws, exceptions):
    _title(ws, "Exceptions", "Items needing attention before sign-off. High = must resolve; Medium = review; "
                             "Info = confirm against the screenshot.")
    heads = ["Severity", "Client", "Form", "Document", "Field / Item", "Issue", "Source", "Drake", "Note", "When",
             "Resolved By", "Resolution"]
    _header(ws, 4, heads, (9, 14, 7, 11, 34, 40, 16, 16, 30, 18, 11, 30))
    order = {"High": 0, "Medium": 1, "Info": 2}
    r = 5
    for e in sorted(exceptions, key=lambda x: order.get(x[0], 3)):
        for j, v in enumerate(e, 1):
            x = ws.cell(r, j, "" if v is None else str(v))
            x.font = _f()
            x.border = THIN
            if j == 6:
                x.alignment = Alignment(wrap_text=True)
        ws.cell(r, 11).fill = YELLOW
        ws.cell(r, 12).fill = YELLOW
        r += 1
    if r == 5:
        ws.cell(5, 1, "No exceptions.").font = _f(italic=True)
    ws.conditional_formatting.add(f"A5:A{max(r, 5)}", CellIsRule(operator="equal", formula=['"High"'], fill=RED))
    ws.conditional_formatting.add(f"A5:A{max(r, 5)}", CellIsRule(operator="equal", formula=['"Medium"'], fill=AMBER))
    ws.freeze_panes = "A5"
    _print_setup(ws, "L", r, "1:4")


def _run_log_sheet(ws, state):
    _title(ws, "Run Log", "Every run of the tool. The input SHA-256 proves which version of the input workbook "
                          "was used.")
    heads = ["Run ID", "Mode", "Started", "Ended", "Status", "Summary", "Initials", "Windows User", "Machine",
             "Input File", "Input SHA-256", "Tool Version"]
    _header(ws, 4, heads, (16, 16, 18, 18, 12, 36, 8, 14, 14, 30, 24, 8))
    r = 5
    for rid, run in state.runs.items():
        vals = [rid, run.get("mode"), run.get("started"), run.get("ended"), run.get("status"), run.get("summary", ""),
                run.get("initials"), run.get("user"), run.get("machine"),
                Path(run.get("input_file", "")).name if run.get("input_file") else "",
                (run.get("input_sha256") or "")[:24], run.get("tool_version")]
        for j, v in enumerate(vals, 1):
            x = ws.cell(r, j, "" if v is None else str(v))
            x.font = _f()
            x.border = THIN
        r += 1
    ws.freeze_panes = "A5"
    _print_setup(ws, "L", r, "1:4")


def _integrity_sheet(ws, ok, n, problem, events, session):
    _title(ws, "Audit Integrity", "The audit log is append-only and hash-chained: each entry includes the SHA-256 "
                                  "of the entry before it, so any edit, deletion or reordering is detected.")
    rows = [
        ("Audit log", str(session.paths["audit_log"].name)),
        ("Chain check", "PASSED - no entries altered, removed or reordered" if ok else f"FAILED - {problem}"),
        ("Entries checked", n),
        ("First entry", events[0]["ts"] if events else ""),
        ("Last entry", events[-1]["ts"] if events else ""),
        ("Last entry hash", events[-1]["hash"] if events else ""),
        ("Checked at", time.strftime("%Y-%m-%d %H:%M:%S")),
        ("Re-check anytime", "python run.py audit-verify"),
        ("Retention", "Keep logs/, evidence/ and tracking/ with the client file per your firm's retention policy."),
    ]
    for i, (a, b) in enumerate(rows, 4):
        ws.cell(i, 1, a).font = _f(True)
        c = ws.cell(i, 2, b)
        c.font = _f(True, color="006100" if ok else "9C0006") if a == "Chain check" else _f()
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 90
    _print_setup(ws, "B", 12)


def _summary_sheet(ws, rows, state, ties):
    _title(ws, "Drake Entry Tracker - Summary",
           f"Built {time.strftime('%Y-%m-%d %H:%M')} from the audit log. Counts are formulas that read each client "
           f"tab. Resolve every item on the Exceptions tab before sign-off.")
    preparers = sorted({run.get("initials", "") for run in state.runs.values()} - {""})
    sign = [("Prepared by", ", ".join(preparers)), ("Reviewed by", ""), ("Review date", "")]
    for i, (a, b) in enumerate(sign, 4):
        ws.cell(i, 1, a).font = _f(True)
        c = ws.cell(i, 2, b)
        c.font = _f()
        if not b:
            c.fill = YELLOW
    heads = ["Client", "Name", "Form", "Documents", "Fields Entered", "Open Mismatches", "Incomplete",
             "Not Read Back", "Corrections", "Tie-Out", "Status", "Reviewer", "Date Reviewed"]
    hr = 8
    _header(ws, hr, heads, (14, 22, 30, 11, 11, 12, 11, 11, 11, 14, 24, 10, 13))
    for i, (sheet, ck, name, fcode, title) in enumerate(rows, hr + 1):
        q = _q(sheet)
        tie = ties.get((ck, fcode), "Not checked")
        tie_txt = {"Yes": "Agreed", "No": "DID NOT AGREE"}.get(tie, "Not confirmed")
        vals = [ck, name, title,
                f'=COUNTIFS({q}!H:H,"{fcode}",{q}!I:I,"INSTANCE")',
                f'=COUNTIFS({q}!H:H,"{fcode}",{q}!I:I,"FIELD")',
                f'=COUNTIFS({q}!H:H,"{fcode}",{q}!D:D,"MISMATCH")',
                f'=COUNTIFS({q}!H:H,"{fcode}",{q}!J:J,"Incomplete")',
                f'=COUNTIFS({q}!H:H,"{fcode}",{q}!D:D,"Not read back")',
                f'=COUNTIFS(Corrections!B:B,A{i},Corrections!C:C,"{fcode}")',
                tie_txt,
                f'=IF(F{i}>0,"Open mismatches",IF(G{i}>0,"Incomplete entries",IF(J{i}="DID NOT AGREE",'
                f'"Tie-out failed",IF(J{i}="Not confirmed","Confirm tie-out",'
                f'IF(H{i}>0,"Ready - check screenshots","Ready for review")))))',
                "", ""]
        for j, v in enumerate(vals, 1):
            x = ws.cell(i, j, v)
            x.font = _f()
            x.border = THIN
        ws.cell(i, 12).fill = YELLOW
        ws.cell(i, 13).fill = YELLOW
    last = max(hr + len(rows), hr + 1)
    ws.conditional_formatting.add(f"K{hr + 1}:K{last}", FormulaRule(
        formula=[f'OR(K{hr + 1}="Open mismatches",K{hr + 1}="Incomplete entries",K{hr + 1}="Tie-out failed")'],
        fill=RED))
    ws.conditional_formatting.add(f"K{hr + 1}:K{last}", FormulaRule(formula=[f'LEFT(K{hr + 1},5)="Ready"'],
                                                                    fill=GREEN))
    ws.freeze_panes = ws.cell(hr + 1, 1)
    _print_setup(ws, "M", last, f"1:{hr}")
