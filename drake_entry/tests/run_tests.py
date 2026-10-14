"""
End-to-end tests against a SIMULATED Drake (tests/fake_drake). No real Drake needed.

  python tests/run_tests.py

Copies the tool into a temp folder, then exercises: validation, entry, a keying
error that stops entry, re-run protection against duplicates, verify, correct,
Drake pre-fill handling, an edited EIN in the input, audit-log tamper detection,
and the tracker workbook.
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
PKG = HERE.parent
FAKE = HERE / "fake_drake"
REPO_SRC = PKG.parent / "src"                     # brokerage_1099 package (repo layout)
FIXTURES = HERE / "fixtures" / "intake_client"
SAMPLE_1099 = PKG.parent / "examples" / "sample_package.json"
PASSED = []


def check(name, cond, detail=""):
    print(f"  [{'PASS' if cond else 'FAIL'}] {name}" + (f"  -- {detail}" if detail and not cond else ""))
    if not cond:
        raise SystemExit(f"\nTEST FAILED: {name}")
    PASSED.append(name)


def run(work, *args, ret="default", env=None):
    e = {**os.environ, "PYTHONPATH": os.pathsep.join([str(FAKE), str(REPO_SRC)]),
         "FAKE_DRAKE_STATE": str(work / "fake_state.json"), "FAKE_RETURN": ret, **(env or {})}
    p = subprocess.run([sys.executable, str(work / "run.py"), *args, "--assume-yes", "--initials", "QA"],
                       cwd=work, env=e, capture_output=True, text=True)
    return p.returncode, p.stdout + p.stderr


def press_esc_in_drake(work, ret):
    """What the user is told to do after a stop: press Esc in Drake."""
    e = {**os.environ, "PYTHONPATH": str(FAKE), "FAKE_DRAKE_STATE": str(work / "fake_state.json"), "FAKE_RETURN": ret}
    subprocess.run([sys.executable, "-c", "import pyautogui; pyautogui.press('esc')"], env=e, cwd=work)


def fake_return(work, ret, screen="W2"):
    s = json.load(open(work / "fake_state.json"))
    return s["returns"].get(ret, {}).get(screen, [])


def write_input(work, rows):
    from openpyxl import load_workbook
    path = work / "input" / "Client_Input.xlsx"
    wb = load_workbook(path)
    ws = wb["W2"]
    names = [c.value for c in ws[2]]
    for i in range(4, 60):
        for j in range(1, len(names) + 1):
            ws.cell(i, j).value = None
    for i, r in enumerate(rows, 4):
        for j, n in enumerate(names, 1):
            if n in r:
                ws.cell(i, j).value = r[n]
    wb.save(path)


SAM_T = {"client_ssn": "123-45-6789", "client_name": "Sam Rivera", "ts": "T", "employer_ein": "84-1234567",
         "employer_name": "Front Range Hotels LLC", "employer_street": "1600 Broadway", "employer_city": "Denver",
         "employer_state": "CO", "employer_zip": "80202", "box1_wages": "52,000.00", "box2_fed_wh": "4800",
         "box3_ss_wages": "55000", "box4_ss_tax": "3400", "box5_medicare_wages": "55000",
         "box6_medicare_tax": "797.50", "box12a_code": "D", "box12a_amt": "3000", "box13_retirement": "X",
         "box15_state": "CO", "box15_state_id": "12345678", "box16_state_wages": "52000", "box17_state_tax": "2288"}
SAM_S = {"client_ssn": "123-45-6789", "client_name": "Sam Rivera", "ts": "S", "employer_ein": "84-7654321",
         "employer_name": "Aurora School District", "employer_city": "Aurora", "employer_state": "CO",
         "employer_zip": "80010", "box1_wages": "38500", "box2_fed_wh": "2900", "box3_ss_wages": "38500",
         "box4_ss_tax": "2387", "box5_medicare_wages": "38500", "box6_medicare_tax": "558.25",
         "box15_state": "CO", "box16_state_wages": "38500", "box17_state_tax": "1694"}
LEE = {"client_ssn": "987-65-4321", "client_name": "Lee Chen", "ts": "T", "employer_ein": "20-1112223",
       "employer_name": "Peak Logistics Inc", "employer_city": "Denver", "employer_state": "CO",
       "employer_zip": "80216", "box1_wages": "61000", "box2_fed_wh": "6100", "box3_ss_wages": "61000",
       "box4_ss_tax": "3782", "box5_medicare_wages": "61000", "box6_medicare_tax": "884.50",
       "box15_state": "CO", "box16_state_wages": "61000", "box17_state_tax": "2684"}
KIM = {"client_ssn": "555-11-2222", "client_name": "Kim Patel", "ts": "T", "employer_ein": "33-4445556",
       "employer_name": "Summit Dental PC", "employer_city": "Boulder", "employer_state": "CO",
       "employer_zip": "80302", "box1_wages": "47000", "box2_fed_wh": "3900", "box3_ss_wages": "47000",
       "box4_ss_tax": "2914", "box5_medicare_wages": "47000", "box6_medicare_tax": "681.50"}


def main():
    work = Path(tempfile.mkdtemp(prefix="drake_entry_test_"))
    shutil.copytree(PKG, work, dirs_exist_ok=True, ignore=shutil.ignore_patterns(
        "logs", "tracking", "evidence", "__pycache__", "*.pyc", "fake_state*"))
    st = json.load(open(work / "settings.json"))
    st["delays"] = {k: 0 for k in st["delays"]}
    json.dump(st, open(work / "settings.json", "w"), indent=2)
    run(work, "make-input", "--force")
    print(f"Test folder: {work}\n")

    print("1. Validation")
    bad = dict(LEE, employer_ein="20-111", box12a_code="ZZ", box12a_amt="100")
    write_input(work, [SAM_T, SAM_S, bad])
    code, out = run(work, "validate")
    check("bad EIN and invalid Box 12 code are errors", code == 1 and "must be 9 digits" in out
          and "not a valid code" in out, out)
    write_input(work, [SAM_T, SAM_S, LEE])
    code, out = run(work, "validate")
    check("clean input validates", code == 0, out)
    check("SS-tax warning raised for Sam's T W-2 (3400 vs 3410 expected)", "WARNING" in out and "Box 4" in out, out)

    print("\n2. Enter")
    code, out = run(work, "enter", "--client", "6789", ret="6789")
    inst = fake_return(work, "6789")
    check("two W-2s entered for Sam", code == 0 and len(inst) == 2, out)
    check("values landed in the right fields", inst[0].get("8") == "52000" and inst[0].get("2") == "841234567"
          and inst[1].get("0") == "S", str(inst[0]))
    check("tracker created", (work / "tracking" / "Entry_Tracker.xlsx").exists())

    print("\n3. Keying error stops entry")
    code, out = run(work, "enter", "--client", "4321", ret="4321", env={"FAKE_CORRUPT": "6100"})
    check("entry stopped on mismatch", code == 1 and "STOPPED" in out and "99999" in out, out)
    press_esc_in_drake(work, "4321")

    print("\n4. Re-running enter doesn't duplicate")
    code, out = run(work, "enter", "--client", "4321", ret="4321")
    check("incomplete W-2 skipped, not re-entered", "INCOMPLETE" in out and len(fake_return(work, "4321")) == 1, out)
    code, out = run(work, "enter", "--client", "6789", ret="6789")
    check("completed W-2s skipped", "already in Drake" in out and len(fake_return(work, "6789")) == 2, out)

    print("\n5. Verify then correct")
    code, out = run(work, "verify", "--client", "4321", ret="4321")
    check("verify finds the differences", "difference(s) logged" in out, out)
    code, out = run(work, "correct", "--client", "4321", ret="4321")
    check("correct fixes them", "correction(s) made and verified" in out, out)
    code, out = run(work, "verify", "--client", "4321", ret="4321")
    check("verify is clean afterwards", "Everything matches the source" in out, out)
    check("Drake now has the right Box 2", fake_return(work, "4321")[0].get("9") == "6100")

    print("\n6. Drake pre-fill is replaced and noted")
    write_input(work, [SAM_T, SAM_S, LEE, KIM])
    code, out = run(work, "enter", "--client", "2222", ret="2222", env={"FAKE_AUTOFILL": "2:3=SUMMIT DENTAL"})
    check("pre-filled employer name replaced with source", code == 0 and
          fake_return(work, "2222")[0].get("3") == "Summit Dental PC", out)
    log = (work / "logs" / "audit_log.jsonl").read_text()
    check("pre-fill noted in audit log", "pre-filled" in log)

    print("\n7. EIN corrected in the input after entry")
    lee2 = dict(LEE, employer_ein="20-1112299")
    write_input(work, [SAM_T, SAM_S, lee2, KIM])
    code, out = run(work, "enter", "--client", "4321", ret="4321")
    check("enter refuses (would duplicate)", "NOT entering" in out and len(fake_return(work, "4321")) == 1, out)
    code, out = run(work, "correct", "--client", "4321", ret="4321")
    check("correct matches it and fixes the EIN", "identity fields changed" in out and
          fake_return(work, "4321")[0].get("2") == "201112299", out)
    code, out = run(work, "verify", "--client", "4321", ret="4321")
    check("verify clean with the new EIN", "Everything matches" in out, out)

    print("\n8. Audit log integrity")
    code, out = run(work, "audit-verify")
    check("audit chain intact", code == 0 and "PASSED" in out, out)
    logp = work / "logs" / "audit_log.jsonl"
    lines = logp.read_text().splitlines()
    backup = list(lines)
    lines[5] = lines[5].replace('"QA"', '"XX"', 1) if '"QA"' in lines[5] else lines[5].replace("OK", "OX", 1)
    logp.write_text("\n".join(lines) + "\n")
    code, out = run(work, "audit-verify")
    check("tampering detected", code == 2 and "FAILED" in out, out)
    logp.write_text("\n".join(backup) + "\n")

    print("\n9. Tracker workbook")
    run(work, "build-tracker")
    from openpyxl import load_workbook
    wb = load_workbook(work / "tracking" / "Entry_Tracker.xlsx")
    check("all sheets present", all(s in wb.sheetnames for s in
                                    ["Summary", "Exceptions", "Corrections", "Tie-Out", "Run Log", "Audit Integrity"]),
          str(wb.sheetnames))
    check("one tab per client", sum(1 for s in wb.sheetnames if s[:4].isdigit()) == 3, str(wb.sheetnames))
    cor = [r for r in wb["Corrections"].iter_rows(min_row=5, values_only=True) if r[0]]
    check("corrections recorded with before/after", len(cor) >= 2 and any(r[5] == "99999" for r in cor), str(cor[:2]))

    intake_tests(work)

    print(f"\nAll {len(PASSED)} checks passed.  Test output kept in {work}")
    return work


def build_client_folder(work):
    """A client folder with every kind of document intake has to sort out."""
    folder = work / "Clients" / "Pat Intake"
    shutil.copytree(FIXTURES, folder)
    w2 = (folder / "W2_FrontRange.txt").read_text()
    (folder / "W2_FrontRange_copy.txt").write_text(w2)                       # same file twice
    (folder / "W2_someone_else.txt").write_text(
        w2.replace("222-33-4444", "999-88-7777").replace("84-1234567", "66-7777777"))
    bank = (folder / "Bank_1099INT.txt").read_text()
    (folder / "Bank_1099INT_CORRECTED.txt").write_text(
        bank.replace("2025 Form 1099-INT", "CORRECTED 2025 Form 1099-INT").replace("150.00", "175.00"))
    cons = (folder / "Brokerage_Consolidated_2025.txt").read_text()
    (folder / "Third_Securities.txt").write_text(                            # does not tie
        cons.replace("EXAMPLE BROKERAGE LLC", "THIRD SECURITIES INC").replace("12-3456789", "77-1234567")
            .replace("XXXX-5678", "XXXX-9999").replace("15,150.00    16,400.00", "15,250.00    16,400.00"))
    pkg = json.loads(SAMPLE_1099.read_text())                                 # keyed JSON, noncovered lots
    pkg.update(client="Pat Intake", recipient_tin_last4="222-33-4444")
    pkg["payers"][0].update(name="Second Brokerage LLC", ein="45-6789012")
    (folder / "Second_Brokerage_keyed.json").write_text(json.dumps(pkg, indent=1))
    return folder


def staged(work, sheet, ssn_last4="4444"):
    import pandas as pd
    df = pd.read_excel(work / "input" / "Client_Input.xlsx", sheet_name=sheet, header=1, dtype=str).fillna("")
    return [r for r in df.to_dict("records") if r["client_ssn"].endswith(ssn_last4)]


def intake_tests(work):
    from openpyxl import load_workbook
    print("\n10. Folder intake: classify every document and stage Schedule B / D and W-2 rows")
    folder = build_client_folder(work)
    code, out = run(work, "intake", str(folder.parent), ret="intake")
    check("intake ran", code == 0 and "Intake report" in out, out)
    expect = {"W2_FrontRange.txt": "Prepared", "W2_FrontRange_copy.txt": "Duplicate file - skipped",
              "W2_someone_else.txt": "Refused - another person's document",
              "Bank_1099INT.txt": "Superseded by CORRECTED", "Bank_1099INT_CORRECTED.txt": "Prepared",
              "Brokerage_Consolidated_2025.txt": "Prepared", "Third_Securities.txt": "Refused - did not tie",
              "Second_Brokerage_keyed.json": "Prepared", "Mortgage_1098.txt": "Recognized - not automated",
              "notes.txt": "Unrecognized", "receipt_photo.jpg": "Needs keyed info"}
    log = [json.loads(l) for l in (work / "logs" / "audit_log.jsonl").read_text().splitlines()]
    got = {e["data"]["file"]: e["data"]["status"] for e in log if e["event"] == "intake_document"}
    check("every document classified with the right status", got == expect,
          str({k: (got.get(k), v) for k, v in expect.items() if got.get(k) != v}))
    rows8949 = staged(work, "8949")
    check("Schedule D staged: summary lines + one line per noncovered lot", len(rows8949) == 11,
          str(len(rows8949)))
    summ = [r for r in rows8949 if r["lot_ref"] == "A-SUMMARY" and "EXAMPLE" in r["description"]][0]
    check("Box A summary carries the wash-sale W adjustment", summ["adj_code"] == "W"
          and summ["adj_amount"] == "500.00" and summ["proceeds"] == "15150.00", str(summ))
    ints = staged(work, "INT")
    check("Schedule B staged from the CORRECTED 1099-INT, not the original",
          any(r["box1"] == "175.00" for r in ints) and not any(r["box1"] == "150.00" for r in ints), str(ints))
    check("W-2 staged once, flagged T by SSN match", len(staged(work, "W2")) == 1
          and staged(work, "W2")[0]["ts"] == "T")
    rep = load_workbook(sorted((work / "tracking").glob("Intake_*.xlsx"))[-1])
    check("intake report has documents / fixes / manual tabs",
          rep.sheetnames == ["Documents", "Fix before entry", "Enter manually"], str(rep.sheetnames))
    manual = [r[2] for r in rep["Enter manually"].iter_rows(min_row=2, values_only=True)]
    check("margin interest and Section 1256 listed for manual entry",
          any("Margin" in str(m) for m in manual) and any("1256" in str(m) or "contracts" in str(m)
                                                         for m in manual), str(manual))

    code, out = run(work, "validate", "--client", "4444")
    check("validate stops on unknown term / missing basis", code == 1 and "unknown term" in out
          and "Cost basis is blank" in out, out[-1500:])
    code, out = run(work, "intake", str(folder.parent), ret="intake")
    check("re-running intake doesn't duplicate rows", "Already prepared" in out
          and len(staged(work, "8949")) == 11, out[-800:])

    # the preparer fills in what the client supplied, straight in the input workbook
    path = work / "input" / "Client_Input.xlsx"
    wb = load_workbook(path)
    ws = wb["8949"]
    names = [c.value for c in ws[2]]
    col = {n: j for j, n in enumerate(names, 1) if n}
    for i in range(4, ws.max_row + 1):
        if not str(ws.cell(i, col["client_ssn"]).value or "").endswith("4444"):
            continue
        for name, value in (("box_8949", "E"), ("cost", "1000.00"), ("date_acquired", "01/15/2015")):
            if not ws.cell(i, col[name]).value:
                ws.cell(i, col[name]).value = value
    wb.save(path)
    code, out = run(work, "intake", str(folder.parent), "--replace", "--enter", ret="intake")
    check("--replace keeps the preparer's entries and --enter types everything into Drake",
          code == 0 and "Entering" in out and "entered" in out, out[-1500:])
    check("Drake has the W-2, 11 8949 lines, 2 DIV, 3 INT, 1 OID, 1 99M",
          [len(fake_return(work, "intake", s)) for s in ("W2", "8949", "DIV", "INT", "OID", "99M")]
          == [1, 11, 2, 3, 1, 1],
          str([len(fake_return(work, "intake", s)) for s in ("W2", "8949", "DIV", "INT", "OID", "99M")]))
    code, out = run(work, "verify", "--client", "4444", ret="intake")
    check("verify: every field in Drake matches the staged source", "Everything matches the source" in out,
          out[-1500:])
    wb = load_workbook(work / "tracking" / "Entry_Tracker.xlsx")
    exc = [r for r in wb["Exceptions"].iter_rows(min_row=5, values_only=True) if r[0]]
    check("tracker exceptions: refused documents High, manual entries listed",
          any(r[0] == "High" and "did not tie" in str(r[5]) for r in exc)
          and any("Enter manually" in str(r[5]) for r in exc), str(exc[:5]))
    code, out = run(work, "audit-verify")
    check("audit chain intact after intake", code == 0 and "PASSED" in out, out)


if __name__ == "__main__":
    main()
