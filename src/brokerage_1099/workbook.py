"""Build the Excel preparer workpaper.

Layout follows the brokerage-1099-summary workpaper standard: Summary first, then lot,
security, wash, noncovered, crypto, income and flag tabs. Every Summary figure is a live
link or formula into a detail tab, and every tie-out is a ROUND(...,2) check cell that is
registered so it can be read back after recalculation.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .analysis import OTHER_CATEGORIES, Analysis, income_label
from .models import (
    BOXES,
    DIGITAL_BOXES,
    INCOME_FORMS,
    NONCOVERED_BOXES,
    SECURITY_BOXES,
    Lot,
    Package,
)

FONT_NAME = "Arial"
BLUE, BLACK, GREEN, WHITE = "0000FF", "000000", "008000", "FFFFFF"
NUM = '#,##0.00_);(#,##0.00);"-"_)'
CHECK_NUM = "#,##0.00_);(#,##0.00);0.00_)"
QTY = "#,##0.0000_);(#,##0.0000)"
DATE_FMT = "mm/dd/yyyy"
PCT = "0.00%"
FILL_HEADER = "D9E1F2"
FILL_SECTION = "1F3864"
FILL_CHECK = "C6EFCE"
FILL_ORANGE = "F8CBAD"
FILL_YELLOW = "FFFF00"

TAB_SUMMARY = "Summary"
TAB_LOTS = "1099-B Lots"
TAB_BYSEC = "By Security"
TAB_WASH = "Wash Sales"
TAB_NONCOV = "Noncovered Lots"
TAB_CRYPTO = "Crypto 1099-DA"
TAB_INCOME = "Income Detail"
TAB_FLAGS = "Flags & Questions"

# Lot table columns (shared by the 1099-B Lots and Crypto 1099-DA tabs)
LOT_HEADERS = [
    ("Box", 6), ("Term", 10), ("Ord", 5), ("Security", 36), ("CUSIP", 12), ("Qty", 12),
    ("Date acquired", 12), ("Date sold", 12), ("Proceeds (d)", 15), ("Cost basis (e)", 15),
    ("Market discount", 13), ("Wash sale (g)", 14), ("Broker gain/(loss)", 15),
    ("Calc gain/(loss)", 15), ("Diff", 10), ("Account", 10), ("Notes", 40), ("Src page", 8),
]
C = {name: get_column_letter(i + 1) for i, name in enumerate(
    ["box", "term", "ord", "security", "cusip", "qty", "acq", "sold", "proceeds", "cost",
     "mkt", "wash", "bgain", "cgain", "diff", "account", "notes", "page"])}

LOT_FIRST_ROW = 5


def q(sheet: str) -> str:
    return "'" + sheet.replace("'", "''") + "'"


def ref(sheet: str, cell: str) -> str:
    return f"{q(sheet)}!{cell}"


def _num(v) -> Optional[float]:
    if v is None:
        return None
    return float(v) if isinstance(v, Decimal) else v


def safe_filename(*parts) -> str:
    raw = "_".join(str(p) for p in parts if p not in (None, ""))
    return re.sub(r"[^A-Za-z0-9_.-]+", "", raw.replace(" ", ""))


@dataclass
class CheckCell:
    label: str
    sheet: str
    cell: str


@dataclass
class LotTable:
    sheet: str
    first: int
    last: int
    box_rows: dict = field(default_factory=dict)  # box -> row in totals block
    broker_rows: dict = field(default_factory=dict)  # box -> row in broker block
    total_row: int = 0
    entry_rows: dict = field(default_factory=dict)  # (box, ordinary) -> row

    def col(self, key: str) -> str:
        c = C[key]
        return f"{q(self.sheet)}!${c}${self.first}:${c}${self.last}"


@dataclass
class BuildResult:
    path: Path
    checks: list[CheckCell]
    status_cell: CheckCell


class _Sheet:
    """Thin styling helper around an openpyxl worksheet."""

    def __init__(self, ws: Worksheet):
        self.ws = ws
        ws.sheet_view.showGridLines = False

    def put(self, row, col, value, kind="text", fmt=None, bold=False, fill=None,
            italic=False, wrap=False, size=10, color=None):
        if isinstance(col, str):
            cell = self.ws[f"{col}{row}"]
        else:
            cell = self.ws.cell(row=row, column=col)
        if isinstance(value, Decimal):
            value = float(value)
        cell.value = value
        if kind in ("text", "input") and isinstance(value, str) and value.startswith("="):
            cell.data_type = "s"  # a label, never a formula
        font_color = color or {"input": BLUE, "link": GREEN}.get(kind, BLACK)
        cell.font = Font(name=FONT_NAME, size=size, bold=bold, italic=italic, color=font_color)
        if kind == "check":
            fill = fill or FILL_CHECK
            fmt = fmt or CHECK_NUM
        elif kind in ("formula", "link"):
            fmt = fmt or NUM
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        if fmt:
            cell.number_format = fmt
        elif isinstance(value, (int, float)) and not isinstance(value, bool):
            cell.number_format = NUM
        elif isinstance(value, date):
            cell.number_format = DATE_FMT
        if wrap:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        return cell

    def header(self, row, labels, start_col=1, fill=FILL_HEADER):
        thin = Side(style="thin", color="8EA9DB")
        for i, label in enumerate(labels):
            if label is None:
                continue
            c = self.put(row, start_col + i, label, bold=True, fill=fill)
            c.alignment = Alignment(wrap_text=True, vertical="center",
                                    horizontal="center" if i else "left")
            c.border = Border(bottom=thin)

    def section(self, row, text, ncols=9, fill=FILL_SECTION, color=WHITE):
        for col in range(1, ncols + 1):
            self.put(row, col, None, fill=fill)
        self.put(row, 1, text, bold=True, fill=fill, color=color, size=11)

    def title(self, title, subtitle):
        self.put(1, 1, title, bold=True, size=14)
        self.put(2, 1, subtitle, italic=True, size=9)

    def widths(self, widths):
        for i, w in enumerate(widths, 1):
            self.ws.column_dimensions[get_column_letter(i)].width = w

    def landscape(self):
        self.ws.page_setup.orientation = "landscape"
        self.ws.page_setup.fitToWidth = 1
        self.ws.page_setup.fitToHeight = 0
        self.ws.sheet_properties.pageSetUpPr.fitToPage = True


class WorkbookBuilder:
    def __init__(self, pkg: Package, analysis: Analysis):
        self.pkg = pkg
        self.an = analysis
        self.wb = Workbook()
        self.checks: list[CheckCell] = []
        self.lot_loc: dict[int, tuple[str, int]] = {}
        self.tables: dict[str, LotTable] = {}
        self.wash_box_rows: dict[str, int] = {}
        self.wash_total_row = 0
        self.income_rows: dict[tuple[str, str], int] = {}
        self.other_block: tuple[int, int] = (0, 0)
        self.state_total_cell = ""

    # -- helpers --------------------------------------------------------------
    def check(self, label, sheet, cell):
        self.checks.append(CheckCell(label, sheet, cell))

    def table_for_box(self, box: str) -> Optional[LotTable]:
        name = TAB_CRYPTO if box in DIGITAL_BOXES else TAB_LOTS
        return self.tables.get(name)

    @property
    def has_digital(self) -> bool:
        return bool(self.pkg.digital_lots) or any(
            t.box in DIGITAL_BOXES for t in self.pkg.box_totals)

    @property
    def has_noncovered(self) -> bool:
        return bool(self.an.noncovered)

    def _lot_boxes(self, digital: bool) -> list[str]:
        candidates = DIGITAL_BOXES if digital else SECURITY_BOXES
        return [b for b in candidates if b in ("A", "D", "G", "J")
                or self.pkg.lots_in([b]) or self.pkg.box_total(b)]

    # -- build ----------------------------------------------------------------
    def build(self, path) -> BuildResult:
        wb = self.wb
        summary = wb.active
        summary.title = TAB_SUMMARY
        order = [TAB_LOTS, TAB_BYSEC, TAB_WASH]
        if self.has_noncovered:
            order.append(TAB_NONCOV)
        if self.has_digital:
            order.append(TAB_CRYPTO)
        order += [TAB_INCOME, TAB_FLAGS]
        for name in order:
            wb.create_sheet(name)

        self._lot_table(TAB_LOTS, self.pkg.security_lots, self._lot_boxes(False),
                        "1099-B Lots", by_asset=False)
        if self.has_digital:
            self._lot_table(TAB_CRYPTO, self.pkg.digital_lots, self._lot_boxes(True),
                            "Crypto 1099-DA Lots", by_asset=True)
            wb[TAB_CRYPTO].sheet_properties.tabColor = "FFC000"
        self._by_security()
        self._wash_sales()
        if self.has_noncovered:
            self._noncovered()
        self._income()
        self._flags()
        status = self._summary()

        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        return BuildResult(path=path, checks=list(self.checks), status_cell=status)

    # -- lot tables -------------------------------------------------------------
    def _lot_table(self, name, lots: list[Lot], boxes, title, by_asset):
        s = _Sheet(self.wb[name])
        p = self.pkg
        s.title(f"{title} - {p.client} - {p.tax_year} {p.broker}",
                "Blue = keyed from the 1099 (page in Src page column). Black = formula. "
                "Calc gain = proceeds - cost + wash; Diff = calc - broker (must be 0).")
        s.widths([w for _, w in LOT_HEADERS])
        s.header(4, [h for h, _ in LOT_HEADERS])
        first = LOT_FIRST_ROW
        lots = sorted(lots, key=lambda l: (list(BOXES).index(l.box), l.description, l.cusip,
                                           str(l.date_sold)))
        r = first
        for lot in lots:
            self.lot_loc[id(lot)] = (name, r)
            s.put(r, C["box"], lot.box, "input")
            s.put(r, C["term"], BOXES[lot.box].term)
            s.put(r, C["ord"], "Y" if lot.ordinary else "N", "input")
            s.put(r, C["security"], lot.description, "input")
            s.put(r, C["cusip"], lot.cusip, "input")
            if lot.quantity is not None:
                s.put(r, C["qty"], lot.quantity, "input", fmt=QTY)
            s.put(r, C["acq"], lot.date_acquired or "N/A", "input")
            s.put(r, C["sold"], lot.date_sold or "N/A", "input")
            s.put(r, C["proceeds"], lot.proceeds, "input")
            s.put(r, C["cost"], lot.cost if lot.cost is not None else "N/A", "input")
            s.put(r, C["mkt"], lot.market_discount, "input")
            s.put(r, C["wash"], lot.wash_sale, "input")
            s.put(r, C["bgain"], lot.broker_gain if lot.broker_gain is not None else "N/A",
                  "input")
            if lot.cost is not None:
                s.put(r, C["cgain"], f"={C['proceeds']}{r}-{C['cost']}{r}+{C['wash']}{r}",
                      "formula", fmt=NUM)
                if lot.broker_gain is not None:
                    s.put(r, C["diff"], f"=ROUND({C['cgain']}{r}-{C['bgain']}{r},2)",
                          "formula", fmt=CHECK_NUM)
            s.put(r, C["account"], lot.account, "input")
            s.put(r, C["notes"], lot.notes, "input")
            if lot.source_page is not None:
                s.put(r, C["page"], lot.source_page, "input", fmt="0")
            if lot.box in NONCOVERED_BOXES:
                for col in ("box", "term", "security"):
                    s.ws[f"{C[col]}{r}"].fill = PatternFill("solid", fgColor=FILL_ORANGE)
            r += 1
        last = max(first, r - 1)
        t = LotTable(name, first, last)
        self.tables[name] = t
        s.ws.freeze_panes = f"A{first}"
        s.ws.auto_filter.ref = f"A4:{C['page']}{last}"
        rng = lambda key: f"${C[key]}${first}:${C[key]}${last}"  # noqa: E731

        # Per-lot check
        r = last + 2
        s.put(r, C["security"], "Per-lot check: sum of |Diff| (must be 0)", bold=True)
        s.put(r, C["diff"], f"=ROUND(SUMPRODUCT(ABS({rng('diff')})),2)", "check")
        self.check(f"{name}: every lot calc = broker gain/(loss)", name, f"{C['diff']}{r}")

        # Totals by box
        r += 2
        s.section(r, "Totals by 8949 box (sum of lots)", ncols=len(LOT_HEADERS))
        r += 1
        s.header(r, ["Box", None, None, "Description", None, "# lots", None, None,
                     "Proceeds (d)", "Cost basis (e)", "Market discount", "Wash sale (g)",
                     "Broker gain/(loss)", "Calc gain/(loss)"])
        r += 1
        top = r
        for box in boxes:
            t.box_rows[box] = r
            s.put(r, "A", box, bold=True)
            s.put(r, "D", BOXES[box].label)
            s.put(r, "F", f"=COUNTIF({rng('box')},$A{r})", "formula", fmt="0")
            for key in ("proceeds", "cost", "mkt", "wash", "bgain", "cgain"):
                s.put(r, C[key], f"=SUMIFS({rng(key)},{rng('box')},$A{r})", "formula", fmt=NUM)
            if box in NONCOVERED_BOXES:
                s.ws[f"A{r}"].fill = PatternFill("solid", fgColor=FILL_ORANGE)
            r += 1
        t.total_row = r
        s.put(r, "A", "Total", bold=True)
        s.put(r, "F", f"=SUM(F{top}:F{r - 1})", "formula", fmt="0", bold=True)
        for key in ("proceeds", "cost", "mkt", "wash", "bgain", "cgain"):
            col = C[key]
            s.put(r, col, f"=SUM({col}{top}:{col}{r - 1})", "formula", bold=True)

        # Broker summary (keyed)
        r += 2
        s.section(r, "Broker 1099-B summary totals (keyed from the summary page)",
                  ncols=len(LOT_HEADERS))
        r += 1
        s.header(r, ["Box", None, None, "Source", None, None, None, None, "Proceeds (d)",
                     "Cost basis (e)", "Market discount", "Wash sale (g)",
                     "Broker gain/(loss)"])
        r += 1
        for box in boxes:
            t.broker_rows[box] = r
            bt = self.pkg.box_total(box)
            s.put(r, "A", box, bold=True)
            if bt is None:
                none_reported = not self.pkg.lots_in([box])
                s.put(r, "D", "None reported" if none_reported else "NOT KEYED - key from "
                      "1099-B summary page", italic=True)
                if none_reported:
                    for key in ("proceeds", "cost", "mkt", "wash", "bgain"):
                        s.put(r, C[key], 0, "input")
            else:
                s.put(r, "D", f"1099-B summary p.{bt.source_page}" if bt.source_page
                      else "1099-B summary page")
                s.put(r, C["proceeds"], bt.proceeds, "input")
                s.put(r, C["cost"], bt.cost if bt.cost is not None else "N/A", "input")
                s.put(r, C["mkt"], bt.market_discount, "input")
                s.put(r, C["wash"], bt.wash_sale, "input")
                s.put(r, C["bgain"], bt.gain if bt.gain is not None else "N/A", "input")
            r += 1

        # Diff block
        r += 1
        s.section(r, "Diff: sum of lots - broker summary (must be 0)", ncols=len(LOT_HEADERS))
        r += 1
        s.header(r, ["Box", None, None, None, None, None, None, None, "Proceeds", "Cost",
                     "Market discount", "Wash sale", "Gain/(loss)"])
        r += 1
        labels = {"proceeds": "proceeds", "cost": "cost", "mkt": "market discount",
                  "wash": "wash", "bgain": "gain/(loss)"}
        for box in boxes:
            s.put(r, "A", box, bold=True)
            br, tr = t.broker_rows[box], t.box_rows[box]
            for key, label in labels.items():
                col = C[key]
                s.put(r, col, f"=ROUND({col}{tr}-N({col}{br}),2)", "check")
                if self.pkg.lots_in([box]) or self.pkg.box_total(box):
                    self.check(f"{name}: Box {box} {label} = broker summary", name,
                               f"{col}{r}")
            r += 1

        # 8949 entry lines (basis reported)
        r += 1
        s.section(r, "Form 8949 entry lines - basis reported (linked to Summary)",
                  ncols=len(LOT_HEADERS))
        r += 1
        s.header(r, ["Box", None, "Ord", "Line", None, "# lots", None, None, "Proceeds (d)",
                     "Cost basis (e)", "Market discount", "Wash sale (g)", "Gain/(loss) (h)"])
        r += 1
        for line in self.an.entry_lines:
            if line.box not in boxes:
                continue
            t.entry_rows[(line.box, line.ordinary)] = r
            s.put(r, "A", line.box, bold=True)
            s.put(r, "C", "Y" if line.ordinary else "N")
            s.put(r, "D", f"{line.term_label} (Box {line.box})")
            crit = f"{rng('box')},$A{r},{rng('ord')},$C{r}"
            s.put(r, "F", f"=COUNTIFS({crit})", "formula", fmt="0")
            for key in ("proceeds", "cost", "mkt", "wash"):
                s.put(r, C[key], f"=SUMIFS({rng(key)},{crit})", "formula")
            s.put(r, C["bgain"], f"={C['proceeds']}{r}-{C['cost']}{r}+{C['wash']}{r}",
                  "formula")
            r += 1

        if by_asset:
            r += 1
            s.section(r, "Subtotals by asset", ncols=len(LOT_HEADERS))
            r += 1
            s.header(r, [None, None, None, "Asset", "CUSIP/ID", "# lots", None, None,
                         "Proceeds (d)", "Cost basis (e)", None, "Wash sale (g)",
                         "Broker gain/(loss)"])
            r += 1
            top = r
            for desc, cusip in _unique_securities(lots):
                key_col, key_val = ("cusip", cusip) if cusip else ("security", desc)
                s.put(r, "D", desc)
                s.put(r, "E", cusip)
                match = f"--({rng(key_col)}={C[key_col]}{r})"
                s.put(r, "F", f"=SUMPRODUCT({match})", "formula", fmt="0")
                for key in ("proceeds", "cost", "wash", "bgain"):
                    s.put(r, C[key], f"=SUMPRODUCT({match},{rng(key)})", "formula")
                r += 1
            s.put(r, "D", "Total", bold=True)
            for key in ("proceeds", "cost", "wash", "bgain"):
                col = C[key]
                s.put(r, col, f"=SUM({col}{top}:{col}{max(top, r - 1)})", "formula", bold=True)
            r += 1
            for key in ("proceeds", "wash"):
                col = C[key]
                s.put(r, col, f"=ROUND({col}{r - 1}-{col}{t.total_row},2)", "check")
                self.check(f"{name}: by-asset {key} = lot total", name, f"{col}{r}")
        s.landscape()

    # -- By Security --------------------------------------------------------------
    def _by_security(self):
        s = _Sheet(self.wb[TAB_BYSEC])
        t = self.tables[TAB_LOTS]
        s.title(f"By Security - {self.pkg.client} - {self.pkg.tax_year}",
                f"All figures are SUMPRODUCTs of the {TAB_LOTS} tab, matched on CUSIP "
                "(or description when no CUSIP).")
        s.widths([38, 12, 8, 15, 15, 14, 16, 16])
        s.header(4, ["Security", "CUSIP", "# lots", "Proceeds", "Basis", "Wash sale",
                     "Net (broker gain/(loss))", "Net (calc)"])
        r = 5
        for desc, cusip in _unique_securities(self.pkg.security_lots):
            key_col = "cusip" if cusip else "security"
            s.put(r, "A", desc)
            s.put(r, "B", cusip)
            local = "B" if cusip else "A"
            match = f"--({t.col(key_col)}=${local}{r})"
            s.put(r, "C", f"=SUMPRODUCT({match})", "formula", fmt="0")
            for col, key in (("D", "proceeds"), ("E", "cost"), ("F", "wash"), ("G", "bgain"),
                             ("H", "cgain")):
                s.put(r, col, f"=SUMPRODUCT({match},{t.col(key)})", "formula")
            r += 1
        last = max(5, r - 1)
        r = max(r, 6)  # with no securities, keep the total off the (empty) data range
        s.put(r, "A", "Total", bold=True)
        s.put(r, "C", f"=SUM(C5:C{last})", "formula", fmt="0", bold=True)
        for col in "DEFGH":
            s.put(r, col, f"=SUM({col}5:{col}{last})", "formula", bold=True)
        tot = r
        r += 1
        s.put(r, "A", f"{TAB_LOTS} total (link)")
        lot_tot = t.total_row
        for col, key in (("C", None), ("D", "proceeds"), ("E", "cost"), ("F", "wash"),
                         ("G", "bgain"), ("H", "cgain")):
            src = f"F{lot_tot}" if key is None else f"{C[key]}{lot_tot}"
            s.put(r, col, f"={ref(TAB_LOTS, src)}", "link", fmt="0" if key is None else NUM)
        r += 1
        s.put(r, "A", "Diff (must be 0)", bold=True)
        for col, label in (("C", "lot count"), ("D", "proceeds"), ("E", "basis"),
                           ("F", "wash"), ("G", "net")):
            s.put(r, col, f"=ROUND({col}{tot}-{col}{r - 1},2)", "check")
            self.check(f"By Security {label} = {TAB_LOTS} total", TAB_BYSEC, f"{col}{r}")
        s.ws.freeze_panes = "A5"
        s.ws.auto_filter.ref = f"A4:H{last}"

    # -- Wash Sales --------------------------------------------------------------
    def _wash_sales(self):
        s = _Sheet(self.wb[TAB_WASH])
        s.title(f"Wash Sales - {self.pkg.client} - {self.pkg.tax_year}",
                "Every lot with a non-zero 1099-B box 1g (keyed off the amount, not the W "
                "code). Green = linked to the lot tab.")
        s.widths([6, 36, 12, 12, 12, 15, 15, 15, 15, 15, 10, 70])
        s.header(4, ["Box", "Security", "CUSIP", "Date acquired", "Date sold", "Proceeds",
                     "Basis", "Disallowed (g)", "Broker gain/(loss)", "Calc gain/(loss)",
                     "Diff", "Notes"])
        wash_lots = [l for l in self.pkg.lots if l.wash_sale]
        wash_lots.sort(key=lambda l: (list(BOXES).index(l.box), l.description,
                                      str(l.date_sold)))
        r = 5
        for lot in wash_lots:
            sheet, lr = self.lot_loc[id(lot)]
            for col, key, fmt in (("A", "box", None), ("B", "security", None),
                                  ("C", "cusip", None), ("D", "acq", DATE_FMT),
                                  ("E", "sold", DATE_FMT), ("F", "proceeds", NUM),
                                  ("G", "cost", NUM), ("H", "wash", NUM), ("I", "bgain", NUM)):
                s.put(r, col, link_formula(sheet, lr, key), "link", fmt=fmt)
            s.put(r, "J", f'=IF(ISNUMBER(G{r}),F{r}-G{r}+H{r},"")', "formula", fmt=NUM)
            s.put(r, "K", f"=IF(AND(ISNUMBER(J{r}),ISNUMBER(I{r})),ROUND(J{r}-I{r},2),0)",
                  "check")
            s.put(r, "L", self.an.wash_notes.get(id(lot), ""), wrap=True, size=9)
            r += 1
        last = max(5, r - 1)
        r = max(r, 6)  # with no wash lines, keep the total off the (empty) data range
        self.wash_range = (5, last)
        s.put(r, "B", "Total", bold=True)
        for col in "FGHI":
            s.put(r, col, f"=SUM({col}5:{col}{last})", "formula", bold=True)
        self.wash_total_row = r
        s.put(r, "K", f"=ROUND(SUMPRODUCT(ABS(K5:K{last})),2)", "check")
        self.check("Wash Sales: every line calc = broker gain/(loss)", TAB_WASH, f"K{r}")

        # Wash total vs lot tab(s)
        r += 1
        s.put(r, "B", "Wash total on lot tab(s) (link)")
        parts = [ref(n, f"{C['wash']}{t.total_row}") for n, t in self.tables.items()]
        s.put(r, "H", "=" + "+".join(parts), "link")
        r += 1
        s.put(r, "B", "Diff: Wash Sales tab - lot tab(s) (must be 0)", bold=True)
        s.put(r, "H", f"=ROUND(H{self.wash_total_row}-H{r - 1},2)", "check")
        self.check("Wash Sales tab total = lot tab wash total", TAB_WASH, f"H{r}")

        # By-security subtotals
        r += 2
        s.section(r, "Wash sales by security", ncols=12)
        r += 1
        s.header(r, ["Box", "Security", "CUSIP", "# lines", None, None, None, "Disallowed"])
        r += 1
        groups = []
        for lot in wash_lots:
            key = (lot.box, lot.description, lot.cusip)
            if key not in groups:
                groups.append(key)
        self.wash_groups = {}
        for box, desc, cusip in groups:
            s.put(r, "A", box)
            s.put(r, "B", desc)
            s.put(r, "C", cusip)
            key_col = "C" if cusip else "B"
            match = (f"--($A$5:$A${last}=$A{r}),--(${key_col}$5:${key_col}${last}"
                     f"=${key_col}{r})")
            s.put(r, "D", f"=SUMPRODUCT({match})", "formula", fmt="0")
            s.put(r, "H", f"=SUMPRODUCT({match},$H$5:$H${last})", "formula")
            self.wash_groups[(box, desc, cusip)] = r
            r += 1

        # Per box: detail vs broker
        r += 1
        s.section(r, "Wash sale totals by box: detail vs broker (must be 0)", ncols=12)
        r += 1
        s.header(r, ["Box", "Description", None, "# lines", None, "Lot tab wash (link)",
                     "Broker wash total (link)", "Wash Sales detail", "Diff vs broker",
                     "Diff vs lot tab"])
        r += 1
        for box in BOXES:
            t = self.table_for_box(box)
            if t is None or box not in t.box_rows:
                continue
            self.wash_box_rows[box] = r
            s.put(r, "A", box, bold=True)
            s.put(r, "B", BOXES[box].label)
            s.put(r, "D", f'=COUNTIF($A$5:$A${last},$A{r})', "formula", fmt="0")
            s.put(r, "F", f"={ref(t.sheet, C['wash'] + str(t.box_rows[box]))}", "link")
            s.put(r, "G", f"=N({ref(t.sheet, C['wash'] + str(t.broker_rows[box]))})", "link")
            s.put(r, "H", f"=SUMIFS($H$5:$H${last},$A$5:$A${last},$A{r})", "formula")
            s.put(r, "I", f"=ROUND(H{r}-G{r},2)", "check")
            s.put(r, "J", f"=ROUND(H{r}-F{r},2)", "check")
            if self.pkg.lots_in([box]) or self.pkg.box_total(box):
                self.check(f"Wash Sales: Box {box} detail = broker wash total", TAB_WASH,
                           f"I{r}")
            r += 1

        r += 1
        s.section(r, "Notes", ncols=12)
        notes = [
            "Reported-basis boxes (A/D/G/J) are entered as one summary line per term with "
            "code W and the disallowed amount in column (g). Not-reported boxes are entered "
            "lot by lot.",
            "The broker tracks only same-account, identical-CUSIP wash sales. Ask about IRAs "
            "(loss permanently disallowed), spouse's accounts, other brokers and options.",
            "Lines where the disallowed amount exceeds the line's own loss reflect broker "
            "aggregation of several lots; they are reported as the broker computed.",
            "Digital assets that are not securities are not subject to the wash sale rule "
            "(confirm no law change for the year). Spot crypto ETFs/trusts are securities.",
        ]
        for n in notes:
            r += 1
            s.put(r, "A", "-")
            s.put(r, "B", n, wrap=False, size=9)
        s.ws.freeze_panes = "A5"
        s.ws.auto_filter.ref = f"A4:L{last}"
        s.landscape()

    # -- Noncovered ----------------------------------------------------------------
    def _noncovered(self):
        ws = self.wb[TAB_NONCOV]
        ws.sheet_properties.tabColor = "ED7D31"
        s = _Sheet(ws)
        s.title(f"Noncovered / Unknown-Term Lots - ENTER EACH LOT INDIVIDUALLY - "
                f"{self.pkg.client} {self.pkg.tax_year}",
                "Green = linked to the lot tab. YELLOW = client must supply (example: date "
                "acquired 03/15/2019, term LT, basis 1,250.00). Basis to use = client basis "
                "if entered, else the 1099 basis.")
        s.widths([6, 8, 34, 12, 11, 12, 12, 14, 14, 13, 14, 14, 9, 14, 14, 14, 34, 8])
        s.header(4, ["Box", "Term", "Security", "CUSIP", "Qty", "Date acquired (1099)",
                     "Date sold", "Proceeds", "Cost (1099)", "Wash sale", "Broker gain/(loss)",
                     "CLIENT date acquired", "CLIENT term (ST/LT)", "CLIENT basis",
                     "Basis to use", "Calc gain/(loss)", "What's missing", "Src page"])
        lots = [l for b in NONCOVERED_BOXES for l in self.an.noncovered.get(b, [])]
        r = 5
        for lot in lots:
            sheet, lr = self.lot_loc[id(lot)]
            for col, key, fmt in (("A", "box", None), ("B", "term", None),
                                  ("C", "security", None), ("D", "cusip", None),
                                  ("E", "qty", QTY), ("F", "acq", DATE_FMT),
                                  ("G", "sold", DATE_FMT), ("H", "proceeds", NUM),
                                  ("I", "cost", NUM), ("J", "wash", NUM), ("K", "bgain", NUM),
                                  ("R", "page", "0")):
                s.put(r, col, link_formula(sheet, lr, key), "link", fmt=fmt)
            for col, fmt in (("L", DATE_FMT), ("M", None), ("N", NUM)):
                s.put(r, col, None, "input", fill=FILL_YELLOW, fmt=fmt)
            s.put(r, "O", f'=IF(ISNUMBER(N{r}),N{r},IF(ISNUMBER(I{r}),I{r},""))', "formula",
                  fmt=NUM)
            s.put(r, "P", f'=IF(ISNUMBER(O{r}),H{r}-O{r}+J{r},"")', "formula", fmt=NUM)
            missing = []
            if lot.box == "X":
                missing.append("term")
            if lot.cost is None:
                missing.append("basis")
            if not isinstance(lot.date_acquired, date):
                missing.append("date acquired")
            s.put(r, "Q", ", ".join(missing) if missing else "verify basis with client",
                  fill=FILL_ORANGE)
            r += 1
        last = max(5, r - 1)
        r = max(r, 6)
        s.put(r, "C", "Total", bold=True)
        for col in "HIJKOP":
            s.put(r, col, f"=SUM({col}5:{col}{last})", "formula", bold=True)

        r += 2
        s.section(r, "Box totals vs broker summary (tie-out only - not keyed as a summary "
                  "line)", ncols=18, fill=FILL_ORANGE, color=BLACK)
        r += 1
        s.header(r, ["Box", None, "Description", "# lots", None, None, None, "Proceeds",
                     "Cost (1099)", "Wash sale", "Broker gain/(loss)", None, None, None,
                     "Broker proceeds (link)", "Broker wash (link)", "Diff proceeds",
                     "Diff wash"])
        r += 1
        for box in NONCOVERED_BOXES:
            t = self.table_for_box(box)
            if box not in self.an.noncovered and not self.pkg.box_total(box):
                continue
            s.put(r, "A", box, bold=True, fill=FILL_ORANGE)
            s.put(r, "C", BOXES[box].label)
            s.put(r, "D", f"=COUNTIF($A$5:$A${last},$A{r})", "formula", fmt="0")
            for col in "HIJK":
                s.put(r, col, f"=SUMIFS({col}$5:{col}${last},$A$5:$A${last},$A{r})", "formula")
            if t is not None and box in t.broker_rows:
                br = t.broker_rows[box]
                s.put(r, "O", f"=N({ref(t.sheet, C['proceeds'] + str(br))})", "link")
                s.put(r, "P", f"=N({ref(t.sheet, C['wash'] + str(br))})", "link")
                s.put(r, "Q", f"=ROUND(H{r}-O{r},2)", "check")
                s.put(r, "R", f"=ROUND(J{r}-P{r},2)", "check")
                self.check(f"Noncovered: Box {box} proceeds = broker summary", TAB_NONCOV,
                           f"Q{r}")
                self.check(f"Noncovered: Box {box} wash = broker summary", TAB_NONCOV, f"R{r}")
            r += 1
        ws.freeze_panes = "A5"
        ws.auto_filter.ref = f"A4:R{last}"
        s.landscape()

    # -- Income ----------------------------------------------------------------------
    def _income(self):
        s = _Sheet(self.wb[TAB_INCOME])
        p = self.pkg
        s.title(f"Income Detail - {p.client} - {p.tax_year} {p.broker}",
                "Payment-level 1099-DIV/INT/OID/MISC lines (blue = keyed). Box totals are "
                "SUMPRODUCTs of the lines; Diff vs the 1099 box amount must be 0.")
        s.widths([8, 7, 40, 38, 14, 14, 15, 10])
        s.header(4, ["Form", "Box", "Description", "Payer / security", "CUSIP", "Pay date",
                     "Amount", "Src page"])

        lines = sorted(p.income_lines,
                       key=lambda l: (_box_sort(l.form, l.box), str(l.pay_date)))
        r = 5
        for line in lines:
            s.put(r, "A", line.form, "input")
            s.put(r, "B", line.box, "input")
            s.put(r, "C", income_label(line.form, line.box)[0])
            s.put(r, "D", line.description, "input")
            s.put(r, "E", line.cusip, "input")
            s.put(r, "F", line.pay_date if line.pay_date else "", "input")
            s.put(r, "G", line.amount, "input")
            if line.source_page is not None:
                s.put(r, "H", line.source_page, "input", fmt="0")
            r += 1
        last = max(5, r - 1)
        self.income_range = (5, last)
        s.ws.auto_filter.ref = f"A4:H{last}"
        s.ws.freeze_panes = "A5"

        # Box totals
        r += 1
        s.section(r, "Totals by box: detail vs 1099 box amount (Diff must be 0)", ncols=8)
        r += 1
        s.header(r, ["Form", "Box", "Description", "Destination", "Detail total",
                     "1099 box amount", "Diff", "Src page"])
        r += 1
        keys = {(l.form, l.box) for l in p.income_lines} | {(t.form, t.box)
                                                            for t in p.form_totals}
        for form, box in sorted(keys, key=lambda k: _box_sort(*k)):
            label, dest = income_label(form, box)
            s.put(r, "A", form, bold=True)
            s.put(r, "B", box, bold=True)
            s.put(r, "C", label)
            s.put(r, "D", dest, size=9)
            s.put(r, "E", f"=SUMPRODUCT(--($A$5:$A${last}=$A{r}),--($B$5:$B${last}=$B{r}),"
                          f"$G$5:$G${last})", "formula")
            ft = p.form_total(form, box)
            if ft is not None:
                s.put(r, "F", ft.amount, "input")
                if ft.source_page is not None:
                    s.put(r, "H", ft.source_page, "input", fmt="0")
            else:
                s.put(r, "F", "NOT KEYED", "input", italic=True)
            s.put(r, "G", f"=ROUND(E{r}-N(F{r}),2)", "check")
            self.check(f"1099-{form} box {box}: detail = box amount", TAB_INCOME, f"G{r}")
            self.income_rows[(form, box)] = r
            r += 1

        # Supplemental
        r += 1
        s.section(r, "Supplemental / non-reportable items and Section 1256", ncols=8)
        r += 1
        s.header(r, ["Category", None, "Description", "Destination", "Amount", None, None,
                     "Src page"])
        r += 1
        top = r
        for item in p.other_items:
            label, dest = OTHER_CATEGORIES.get(item.category, OTHER_CATEGORIES["other"])
            s.put(r, "A", item.category)
            s.put(r, "C", item.description or label, "input")
            s.put(r, "D", dest, size=9)
            s.put(r, "E", item.amount, "input")
            if item.source_page is not None:
                s.put(r, "H", item.source_page, "input", fmt="0")
            r += 1
        self.other_block = (top, max(top, r - 1))

        # State Treasury obligations
        r += 1
        s.section(r, "State: U.S. Treasury / direct federal obligation subtraction", ncols=8)
        r += 1
        s.header(r, ["", None, "Security", "CUSIP", "% from 1099", "Fund ordinary dividends",
                     "State subtraction", "Src page"])
        r += 1
        top = r
        for ob in p.state_obligations:
            s.put(r, "C", ob.description, "input")
            s.put(r, "D", ob.cusip, "input")
            s.put(r, "E", ob.percent, "input", fmt=PCT)
            if ob.dividend_amount is not None:
                s.put(r, "F", ob.dividend_amount, "input")
            else:
                s.put(r, "F", f'=SUMPRODUCT(--($A$5:$A${last}="DIV"),--($B$5:$B${last}="1a"),'
                              f"--($E$5:$E${last}=$D{r}),$G$5:$G${last})", "formula")
            s.put(r, "G", f"=ROUND(E{r}*F{r},2)", "formula")
            if ob.source_page is not None:
                s.put(r, "H", ob.source_page, "input", fmt="0")
            r += 1
        for form, box, text in (("INT", "3", "1099-INT box 3 (Treasury interest)"),
                                ("OID", "8", "1099-OID box 8 (Treasury OID)")):
            if (form, box) in self.income_rows:
                s.put(r, "C", text)
                s.put(r, "G", f"=E{self.income_rows[(form, box)]}", "formula")
                r += 1
        s.put(r, "C", "Total state subtraction", bold=True)
        s.put(r, "G", f"=SUM(G{top}:G{max(top, r - 1)})" if r > top else "=0", "formula",
              bold=True)
        self.state_total_cell = f"G{r}"
        s.landscape()

    # -- Flags -------------------------------------------------------------------------
    def _flags(self):
        s = _Sheet(self.wb[TAB_FLAGS])
        s.title(f"Flags & Questions - {self.pkg.client} - {self.pkg.tax_year}",
                "Amounts are computed by the program from the 1099 data. Update Status as "
                "items are resolved.")
        s.widths([5, 22, 100, 15, 10])
        s.header(4, ["#", "Category", "Issue / question", "Amount", "Status"])
        for i, f in enumerate(self.an.flags, 1):
            r = 4 + i
            s.put(r, "A", i, fmt="0")
            s.put(r, "B", f.category)
            s.put(r, "C", f.issue, wrap=True)
            if f.amount is not None:
                s.put(r, "D", f.amount)
            s.put(r, "E", f.status, fill=FILL_YELLOW if f.status == "Open" else None)
        s.ws.freeze_panes = "A5"

    # -- Summary -----------------------------------------------------------------------
    def _summary(self) -> CheckCell:
        s = _Sheet(self.wb[TAB_SUMMARY])
        p, an = self.pkg, self.an
        s.widths([46, 13, 16, 16, 12, 16, 16, 14, 40])
        s.title(f"Brokerage 1099 Summary - {p.client}",
                "Preparer workpaper for data entry - not tax advice. Recipient TIN and "
                "account numbers shown as last 4 only.")
        r = 4
        hdr = [
            ("Client", p.client),
            ("Tax year", str(p.tax_year)),
            ("Broker", p.broker),
            ("Statement date", p.statement_date or ""),
            ("Status", "CORRECTED" if p.corrected else "Original"),
        ]
        if p.corrected and p.corrected_message:
            hdr.append(("Corrected message", f'"{p.corrected_message}"'))
        for change in p.prior_changes:
            hdr.append(("Change from prior version", change))
        hdr += [
            ("Recipient TIN", p.recipient_tin_last4 or "not recorded"),
            ("Forms present", ", ".join(p.forms_present) or "see payer table"),
        ]
        if p.source_file:
            hdr.append(("Source file", p.source_file))
        for label, value in hdr:
            s.put(r, "A", label, bold=True)
            s.put(r, "B", value, "input", fill=FILL_ORANGE if value == "CORRECTED" else None)
            r += 1
        r += 1
        s.header(r, ["Payer name", "Payer EIN", None, "Account", None, "Forms covered"])
        r += 1
        for payer in p.payers or []:
            s.put(r, "A", payer.name, "input")
            s.put(r, "B", payer.ein, "input")
            s.put(r, "D", payer.account_last4, "input")
            s.put(r, "F", ", ".join(payer.forms), "input")
            r += 1
        if not p.payers:
            s.put(r, "A", "NOT RECORDED - key payer name and EIN", italic=True)
            r += 1

        # 2. 8949 entry
        r += 1
        s.section(r, "Form 8949 one-line entry - basis REPORTED to IRS")
        r += 1
        s.header(r, ["Term", "8949 box", "(d) Proceeds", "(e) Cost basis", "(f) Code",
                     "(g) Wash sale adj", "(h) Gain/(loss)", "Payer EIN", "Note"])
        r += 1
        entry_cells: dict[str, list[int]] = {}
        for line in an.entry_lines:
            t = self.table_for_box(line.box)
            er = t.entry_rows[(line.box, line.ordinary)]
            s.put(r, "A", line.term_label)
            s.put(r, "B", line.box)
            s.put(r, "C", f"={ref(t.sheet, C['proceeds'] + str(er))}", "link")
            s.put(r, "D", f"={ref(t.sheet, C['cost'] + str(er))}", "link")
            s.put(r, "E", f'=IF(F{r}<>0,"W","")', "formula")
            s.put(r, "F", f"={ref(t.sheet, C['wash'] + str(er))}", "link")
            s.put(r, "G", f"=C{r}-D{r}+F{r}", "formula")
            s.put(r, "H", p.payer_ein_for_box(line.box), "input")
            if line.lot_count == 0:
                s.put(r, "I", "None reported", italic=True)
            elif line.market_discount:
                s.put(r, "I", f"Includes market discount {line.market_discount:,.2f} "
                      "(ordinary)", size=9)
            entry_cells.setdefault(line.box, []).append(r)
            r += 1

        # 3. Tie-out per term
        r += 1
        s.section(r, "Tie-out to amounts reported (per term)")
        for box, rows in entry_cells.items():
            r += 1
            t = self.table_for_box(box)
            bt = p.box_total(box)
            s.put(r, "A", f"{BOXES[box].label} (Box {box})", bold=True)
            r += 1
            start = r
            s.put(r, "A", "Proceeds - Cost = Gain/(loss) before wash")
            s.put(r, "G", "=" + "+".join(f"(C{x}-D{x})" for x in rows), "formula")
            r += 1
            s.put(r, "A", "Add: wash sale loss disallowed (Wash Sales tab)")
            s.put(r, "G", f"={ref(TAB_WASH, 'H' + str(self.wash_box_rows[box]))}", "link")
            r += 1
            s.put(r, "A", "Equals: reportable gain/(loss)", bold=True)
            s.put(r, "G", f"=G{start}+G{start + 1}", "formula", bold=True)
            r += 1
            src = f" p.{bt.source_page}" if bt and bt.source_page else ""
            s.put(r, "A", f"Broker-reported gain/(loss) (1099-B summary{src})")
            if bt is not None and bt.gain is not None:
                s.put(r, "G", bt.gain, "input")
            elif not p.lots_in([box]):
                s.put(r, "G", 0, "input")
                s.put(r, "I", "None reported", italic=True)
            else:
                s.put(r, "I", "NOT KEYED - key from the 1099-B summary page", italic=True)
            r += 1
            s.put(r, "A", "Diff (must be 0)", bold=True)
            s.put(r, "G", f"=ROUND(G{r - 2}-N(G{r - 1}),2)", "check")
            self.check(f"Summary: Box {box} proceeds - cost + wash = broker gain/(loss)",
                       TAB_SUMMARY, f"G{r}")
            r += 1
            # wash detail by security
            groups = [(k, v) for k, v in self.wash_groups.items() if k[0] == box]
            if not groups:
                s.put(r, "A", f"No wash sales reported in Box {box}.", italic=True)
                r += 1
                continue
            s.header(r, ["Wash sale detail by security", "CUSIP", None, None, None,
                         "# lines", "Disallowed"])
            r += 1
            top = r
            for (_, desc, cusip), wr in groups:
                s.put(r, "A", desc)
                s.put(r, "B", cusip)
                s.put(r, "F", f"={ref(TAB_WASH, 'D' + str(wr))}", "link", fmt="0")
                s.put(r, "G", f"={ref(TAB_WASH, 'H' + str(wr))}", "link")
                r += 1
            s.put(r, "A", "Total", bold=True)
            s.put(r, "F", f"=SUM(F{top}:F{r - 1})", "formula", fmt="0", bold=True)
            s.put(r, "G", f"=SUM(G{top}:G{r - 1})", "formula", bold=True)
            r += 1
            s.put(r, "A", "Broker wash total (1099-B summary)")
            s.put(r, "G", f"=N({ref(t.sheet, C['wash'] + str(t.broker_rows[box]))})", "link")
            r += 1
            s.put(r, "A", "Diff (must be 0)", bold=True)
            s.put(r, "G", f"=ROUND(G{r - 2}-G{r - 1},2)", "check")
            self.check(f"Summary: Box {box} wash by security = broker wash", TAB_SUMMARY,
                       f"G{r}")
            r += 1

        # 4. Not reported block
        r += 1
        if an.noncovered:
            s.section(r, "Basis NOT reported / noncovered / unknown term - ENTER EACH LOT "
                         "INDIVIDUALLY", fill=FILL_ORANGE, color=BLACK)
            r += 1
            s.header(r, ["Box", None, "Proceeds", "Cost (1099)", "# lots", None, None,
                         "Enter from tab", "Note"], fill=FILL_ORANGE)
            r += 1
            for box, lots in an.noncovered.items():
                t = self.table_for_box(box)
                tr = t.box_rows[box]
                s.put(r, "A", f"Box {box} - {BOXES[box].label}", fill=FILL_ORANGE, bold=True)
                s.put(r, "C", f"={ref(t.sheet, C['proceeds'] + str(tr))}", "link")
                s.put(r, "D", f"={ref(t.sheet, C['cost'] + str(tr))}", "link")
                s.put(r, "E", f"={ref(t.sheet, 'F' + str(tr))}", "link", fmt="0")
                s.put(r, "H", TAB_NONCOV)
                na = sum(1 for l in lots if l.cost is None)
                note = f"Cost N/A on {na} lot(s)" if na else ""
                if box == "X":
                    note = ("Term AND basis must come from the client. " + note).strip()
                s.put(r, "I", note, size=9)
                r += 1
            s.put(r, "A", "Totals above are for tie-out only and are NOT keyed as a summary "
                          "line.", italic=True, size=9)
            r += 1
            all_nc = [l for lots in an.noncovered.values() for l in lots]
            if len(all_nc) <= 15:
                s.header(r, ["Security", "CUSIP", "Qty", "Date acquired", "Date sold",
                             "Proceeds", "Cost", "Box", "Missing"], fill=FILL_ORANGE)
                r += 1
                for lot in all_nc:
                    sheet, lr = self.lot_loc[id(lot)]
                    for col, key, fmt in (("A", "security", None), ("B", "cusip", None),
                                          ("C", "qty", QTY), ("D", "acq", DATE_FMT),
                                          ("E", "sold", DATE_FMT), ("F", "proceeds", NUM),
                                          ("G", "cost", NUM), ("H", "box", None)):
                        s.put(r, col, link_formula(sheet, lr, key), "link", fmt=fmt)
                    missing = (["term"] if lot.box == "X" else []) + \
                        (["basis"] if lot.cost is None else []) + \
                        ([] if isinstance(lot.date_acquired, date) else ["date acquired"])
                    s.put(r, "I", ", ".join(missing) or "verify basis", fill=FILL_YELLOW
                          if missing else None)
                    r += 1
        else:
            s.section(r, "Basis NOT reported / noncovered / unknown term")
            r += 1
            s.put(r, "A", "None reported (Boxes B, E, X, H, K are all zero).", italic=True)
            r += 1

        # 5. Income
        r += 1
        s.section(r, "Income headline items")
        r += 1
        s.header(r, ["Item", "Form / box", "Amount", "Destination", None, None, None,
                     "Source tab"])
        r += 1
        totals = {}
        for line in p.income_lines:
            totals[(line.form, line.box)] = totals.get((line.form, line.box), 0) + line.amount
        zero_forms = set()
        for key, row in sorted(self.income_rows.items(),
                               key=lambda kv: kv[1]):
            form, box = key
            ft = p.form_total(form, box)
            amount = ft.amount if ft else totals.get(key, 0)
            if not amount:
                zero_forms.add(form)
                continue
            label, dest = income_label(form, box)
            s.put(r, "A", label)
            s.put(r, "B", f"{form} {box}")
            s.put(r, "C", f"={ref(TAB_INCOME, 'E' + str(row))}", "link")
            s.put(r, "D", dest, size=9)
            s.put(r, "H", TAB_INCOME)
            r += 1
        present = sorted({f for f, _ in self.income_rows} | zero_forms)
        s.put(r, "A", "All other " + ", ".join(f"1099-{f}" for f in present or INCOME_FORMS)
              + " boxes: 0", italic=True)
        r += 1
        other_top, other_bot = self.other_block
        cats = []
        for item in p.other_items:
            if item.category not in cats:
                cats.append(item.category)
        for cat in cats:
            label, dest = OTHER_CATEGORIES.get(cat, OTHER_CATEGORIES["other"])
            s.put(r, "A", label)
            s.put(r, "B", "Supplemental")
            s.put(r, "C", f'=SUMIFS({ref(TAB_INCOME, f"$E${other_top}:$E${other_bot}")},'
                          f'{ref(TAB_INCOME, f"$A${other_top}:$A${other_bot}")},"{cat}")',
                  "link")
            s.put(r, "D", dest, size=9)
            s.put(r, "H", TAB_INCOME)
            r += 1
        s.put(r, "A", "State: U.S. Treasury / federal obligation subtraction")
        s.put(r, "B", "State")
        s.put(r, "C", f"={ref(TAB_INCOME, self.state_total_cell)}", "link")
        s.put(r, "D", "State subtraction (e.g. Colorado)", size=9)
        s.put(r, "H", TAB_INCOME)
        r += 1
        if an.digital_asset_question_yes:
            s.put(r, "A", "Form 1040 digital asset question", bold=True)
            s.put(r, "C", "YES", bold=True)
            r += 1

        # 6. Checks
        r += 1
        s.section(r, "Tie-out checks (every value must be 0.00)")
        r += 1
        s.header(r, ["Check", None, "Value", "Location"])
        r += 1
        top = r
        for chk in list(self.checks):
            s.put(r, "A", chk.label, size=9)
            target = chk.cell if chk.sheet == TAB_SUMMARY else ref(chk.sheet, chk.cell)
            s.put(r, "C", f"={target}", "link" if chk.sheet != TAB_SUMMARY else "formula",
                  fmt=CHECK_NUM, fill=FILL_CHECK)
            s.put(r, "D", f"{chk.sheet}!{chk.cell}", size=9)
            r += 1
        bot = max(top, r - 1)
        r += 1
        s.put(r, "A", "OVERALL STATUS", bold=True, size=12)
        s.put(r, "C", f'=IF(SUMPRODUCT(--(ABS(C{top}:C{bot})>=0.005))=0,"ALL TIED","CHECK")',
              "formula", bold=True, size=12, fill=FILL_CHECK)
        status = CheckCell("Overall status", TAB_SUMMARY, f"C{r}")

        # 7. Legend
        r += 2
        s.section(r, "Color legend")
        legend = [
            ("Blue text = keyed from the 1099 (source page noted)", "input", None),
            ("Black text = formula", "formula", None),
            ("Green text = link to another tab", "link", None),
            ("Green fill = check cell (must be 0.00)", "text", FILL_CHECK),
            ("Orange fill = noncovered / unknown term / enter individually", "text",
             FILL_ORANGE),
            ("Yellow fill = client must supply", "text", FILL_YELLOW),
        ]
        for text, kind, fill in legend:
            r += 1
            s.put(r, "A", text, kind, fill=fill)
        s.landscape()
        return status


TEXT_KEYS = {"box", "term", "ord", "security", "cusip", "account", "notes"}


def link_formula(sheet: str, row: int, key: str) -> str:
    """Formula linking to a lot-table cell. Text fields go through T() so an empty
    source shows blank rather than 0."""
    target = ref(sheet, f"{C[key]}{row}")
    return f"=T({target})" if key in TEXT_KEYS else f"={target}"


def _box_sort(form: str, box: str):
    m = re.match(r"(\d+)([a-z]?)", box)
    n = (int(m.group(1)), m.group(2)) if m else (999, box)
    return (INCOME_FORMS.index(form) if form in INCOME_FORMS else 99), n


def _unique_securities(lots) -> list[tuple[str, str]]:
    seen, out = set(), []
    for lot in sorted(lots, key=lambda l: (l.description, l.cusip)):
        key = lot.cusip or lot.description
        if key in seen:
            continue
        seen.add(key)
        out.append((lot.description, lot.cusip))
    return out


def default_filename(pkg: Package) -> str:
    return safe_filename(pkg.client, pkg.tax_year, pkg.broker, "1099", "Summary") + ".xlsx"


def build_workbook(pkg: Package, analysis: Analysis, path) -> BuildResult:
    return WorkbookBuilder(pkg, analysis).build(path)
