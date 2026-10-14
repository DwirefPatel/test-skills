"""Recalculate a workbook with headless LibreOffice and read the check cells back.

openpyxl writes formulas without cached values; LibreOffice computes them so the file
opens with numbers in any viewer and so the check cells can be verified here.
"""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from .workbook import BuildResult

EXCEL_ERRORS = ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A")
TOLERANCE = 0.005


@dataclass
class RecalcResult:
    ok: bool
    error: str = ""
    formula_errors: list[str] = field(default_factory=list)


@dataclass
class CheckReadback:
    label: str
    location: str
    value: object

    @property
    def ok(self) -> bool:
        return isinstance(self.value, (int, float)) and abs(self.value) < TOLERANCE


@dataclass
class VerifyResult:
    recalc: RecalcResult
    checks: list[CheckReadback] = field(default_factory=list)
    status: object = None

    @property
    def failures(self) -> list[CheckReadback]:
        return [c for c in self.checks if not c.ok]

    @property
    def ok(self) -> bool:
        return (self.recalc.ok and not self.recalc.formula_errors and not self.failures
                and self.status == "ALL TIED")


def soffice_binary() -> str | None:
    return shutil.which("soffice") or shutil.which("libreoffice")


def recalc_in_place(path, timeout: int = 180) -> RecalcResult:
    path = Path(path).resolve()
    binary = soffice_binary()
    if not binary:
        return RecalcResult(False, "LibreOffice (soffice) not found; workbook not recalculated")
    with tempfile.TemporaryDirectory(prefix="b1099-lo-") as tmp:
        tmp = Path(tmp)
        out_dir = tmp / "out"
        cmd = [binary, f"-env:UserInstallation={(tmp / 'profile').as_uri()}", "--headless",
               "--norestore", "--calc", "--convert-to", "xlsx:Calc Office Open XML",
               "--outdir", str(out_dir), str(path)]
        try:
            proc = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
        except subprocess.TimeoutExpired:
            return RecalcResult(False, f"LibreOffice timed out after {timeout}s")
        out = out_dir / path.name
        if proc.returncode != 0 or not out.exists():
            return RecalcResult(False, f"LibreOffice failed: {proc.stderr.strip()[:500]}")
        shutil.copyfile(out, path)

    errors = []
    wb = load_workbook(path, data_only=True)
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value in EXCEL_ERRORS:
                        errors.append(f"{ws.title}!{cell.coordinate} {cell.value}")
    finally:
        wb.close()
    return RecalcResult(True, formula_errors=errors)


def read_checks(build: BuildResult) -> tuple[list[CheckReadback], object]:
    wb = load_workbook(build.path, data_only=True)
    try:
        checks = [CheckReadback(c.label, f"{c.sheet}!{c.cell}", wb[c.sheet][c.cell].value)
                  for c in build.checks]
        status = wb[build.status_cell.sheet][build.status_cell.cell].value
    finally:
        wb.close()
    return checks, status


def recalc_and_verify(build: BuildResult, timeout: int = 180) -> VerifyResult:
    rc = recalc_in_place(build.path, timeout)
    if not rc.ok:
        return VerifyResult(rc)
    checks, status = read_checks(build)
    return VerifyResult(rc, checks, status)
