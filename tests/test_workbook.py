from openpyxl import load_workbook

from brokerage_1099 import package_from_dict, process_package
from brokerage_1099.cli import main

from .conftest import SAMPLE, needs_soffice

EXPECTED_TABS = ["Summary", "1099-B Lots", "By Security", "Wash Sales", "Noncovered Lots",
                 "Crypto 1099-DA", "Income Detail", "Flags & Questions"]


def test_build_without_recalc(sample_pkg, tmp_path):
    result = process_package(sample_pkg, tmp_path, recalc=False)
    assert result.workbook_path.name == "SampleClient_2025_ExampleBrokerage_1099_Summary.xlsx"
    wb = load_workbook(result.workbook_path)
    assert wb.sheetnames == EXPECTED_TABS
    ws = wb["Summary"]
    # Summary entry values are links into detail tabs, never re-keyed numbers
    entry = [r for r in ws.iter_rows(values_only=True) if r[0] == "Short-term"][0]
    assert str(entry[2]).startswith("='1099-B Lots'!")
    assert entry[7] == "12-3456789"
    da = [r for r in ws.iter_rows(values_only=True) if r[0] == "Digital asset short-term"][0]
    assert da[7] == "98-7654321"
    for name in wb.sheetnames:
        assert wb[name].sheet_view.showGridLines is False
    font_names = {c.font.name for row in ws.iter_rows() for c in row if c.value is not None}
    assert font_names == {"Arial"}
    assert result.validation.ok


def test_tabs_omitted_when_not_needed(sample_dict, tmp_path):
    sample_dict["lots"] = [l for l in sample_dict["lots"] if l["box"] in ("A", "D")]
    sample_dict["box_totals"] = [t for t in sample_dict["box_totals"] if t["box"] in ("A", "D")]
    result = process_package(package_from_dict(sample_dict), tmp_path, recalc=False)
    names = load_workbook(result.workbook_path).sheetnames
    assert "Noncovered Lots" not in names and "Crypto 1099-DA" not in names


@needs_soffice
def test_recalc_all_tied(sample_pkg, tmp_path):
    result = process_package(sample_pkg, tmp_path)
    assert result.verify.recalc.ok, result.verify.recalc.error
    assert result.verify.recalc.formula_errors == []
    assert result.verify.status == "ALL TIED"
    assert result.ok
    assert "ALL TIED" in result.headline


@needs_soffice
def test_recalc_reports_break_location(sample_dict, tmp_path):
    sample_dict["lots"][0]["broker_gain"] = "-700.00"
    result = process_package(package_from_dict(sample_dict), tmp_path)
    assert result.verify.status == "CHECK"
    locations = {c.location for c in result.verify.failures}
    assert any(loc.startswith("1099-B Lots!") for loc in locations)
    assert not result.ok


def test_cli_build_no_recalc(tmp_path, capsys):
    code = main(["build", str(SAMPLE), "-o", str(tmp_path), "--no-recalc",
                 "--report", str(tmp_path / "r.json")])
    assert code == 0
    assert (tmp_path / "r.json").exists()
    assert "Workbook:" in capsys.readouterr().out
